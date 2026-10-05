"""Reviewed server-owned DOCX templates and bounded field rendering."""

from __future__ import annotations

import hashlib
import io
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from docx import Document
from docx.document import Document as DocumentObject
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from pydantic import BaseModel, ConfigDict, Field, field_validator

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PROJECT_BRIEF_TEMPLATE_ID = "project-brief-v1"
TEMPLATE_OPERATION = "template_to_docx"
TEMPLATE_WORKER_VERSION = "audrey-template-worker/1"
_MAX_DOCX_BYTES = 5 * 1024 * 1024
_MAX_PACKAGE_FILES = 128
_MAX_PACKAGE_UNCOMPRESSED_BYTES = 12 * 1024 * 1024
_PLACEHOLDER_RE = re.compile(r"\{\{([a-z][a-z0-9_]*)\}\}")
_EXTERNAL_RELATIONSHIP_RE = re.compile(
    rb"TargetMode\s*=\s*(['\"])\s*External\s*\1",
    re.IGNORECASE,
)
_PROJECT_BRIEF_PLACEHOLDERS = frozenset({
    "title",
    "prepared_for",
    "prepared_on",
    "summary",
    "objectives",
    "next_steps",
})


class DocumentTemplateError(ValueError):
    """A template, field set, or rendered package violated the fixed contract."""


class ProjectBriefFields(BaseModel):
    """Plain-text fields accepted by the first reviewed document template."""

    model_config = ConfigDict(extra="forbid", str_strip_whitespace=True)

    title: str = Field(min_length=1, max_length=120)
    prepared_for: str = Field(min_length=1, max_length=120)
    prepared_on: str = Field(min_length=1, max_length=40)
    summary: str = Field(min_length=1, max_length=2_000)
    objectives: list[str] = Field(min_length=1, max_length=8)
    next_steps: list[str] = Field(min_length=1, max_length=8)

    @field_validator("title", "prepared_for", "prepared_on", "summary")
    @classmethod
    def validate_plain_text(cls, value: str) -> str:
        return _plain_text(value, maximum=max(2_000, len(value)))

    @field_validator("objectives", "next_steps")
    @classmethod
    def validate_plain_text_items(cls, values: list[str]) -> list[str]:
        return [_plain_text(value, maximum=300) for value in values]


class TemplateDocumentArguments(BaseModel):
    """Normalized arguments stored behind an exact approval digest."""

    model_config = ConfigDict(extra="forbid", str_strip_whitespace=True)

    template_id: str = Field(min_length=1, max_length=100)
    template_sha256: str = Field(min_length=64, max_length=64, pattern=r"^[0-9a-f]{64}$")
    filename: str = Field(min_length=1, max_length=255)
    fields: ProjectBriefFields

    @field_validator("filename")
    @classmethod
    def validate_filename(cls, value: str) -> str:
        cleaned = value.strip()
        if Path(cleaned).name != cleaned or not cleaned.casefold().endswith(".docx"):
            raise ValueError("filename must be a plain .docx name")
        if any(ord(char) < 32 for char in cleaned):
            raise ValueError("filename contains control characters")
        return cleaned


@dataclass(frozen=True, slots=True)
class TemplateDescriptor:
    template_id: str
    name: str
    description: str
    version: int
    content_sha256: str
    bytes_count: int
    fields: tuple[str, ...]


@dataclass(frozen=True, slots=True)
class RenderedDocument:
    content: bytes
    content_sha256: str
    text: str
    expected_text: tuple[str, ...]


class DocumentTemplateCatalog:
    """Load, audit, and render Audrey's fixed DOCX template catalogue."""

    def __init__(self, asset_root: Path | None = None) -> None:
        self._asset_root = asset_root or Path(__file__).with_name("assets")
        self._template_path = self._asset_root / "project-brief-v1.docx"
        try:
            self._template_bytes = self._template_path.read_bytes()
        except OSError as exc:
            raise DocumentTemplateError("project brief template is unavailable") from exc
        self._template_sha256 = hashlib.sha256(self._template_bytes).hexdigest()
        _validate_package(self._template_bytes, require_placeholders=True)
        template = _open_document(self._template_bytes)
        placeholders = _placeholders(template)
        if placeholders != _PROJECT_BRIEF_PLACEHOLDERS:
            raise DocumentTemplateError(
                "project brief template placeholders do not match its typed schema"
            )

    def list(self) -> tuple[TemplateDescriptor, ...]:
        return (self.descriptor(PROJECT_BRIEF_TEMPLATE_ID),)

    def descriptor(self, template_id: str) -> TemplateDescriptor:
        if template_id != PROJECT_BRIEF_TEMPLATE_ID:
            raise DocumentTemplateError("unknown document template")
        return TemplateDescriptor(
            template_id=PROJECT_BRIEF_TEMPLATE_ID,
            name="Project brief",
            description=(
                "Create a polished brief with a title, recipient, summary, objectives, "
                "and next steps."
            ),
            version=1,
            content_sha256=self._template_sha256,
            bytes_count=len(self._template_bytes),
            fields=(
                "title",
                "prepared_for",
                "prepared_on",
                "summary",
                "objectives",
                "next_steps",
            ),
        )

    def validate_arguments(self, value: dict[str, Any]) -> TemplateDocumentArguments:
        try:
            arguments = TemplateDocumentArguments.model_validate(value)
        except Exception as exc:
            raise DocumentTemplateError("template arguments are invalid") from exc
        descriptor = self.descriptor(arguments.template_id)
        if arguments.template_sha256 != descriptor.content_sha256:
            raise DocumentTemplateError("approved template digest does not match the server template")
        return arguments

    def render(self, value: TemplateDocumentArguments | dict[str, Any]) -> RenderedDocument:
        arguments = (
            value if isinstance(value, TemplateDocumentArguments)
            else self.validate_arguments(value)
        )
        descriptor = self.descriptor(arguments.template_id)
        if arguments.template_sha256 != descriptor.content_sha256:
            raise DocumentTemplateError("approved template digest does not match the server template")
        fields = arguments.fields
        replacements = {
            "title": fields.title,
            "prepared_for": fields.prepared_for,
            "prepared_on": fields.prepared_on,
            "summary": fields.summary,
            "objectives": "\n".join(f"• {item}" for item in fields.objectives),
            "next_steps": "\n".join(f"• {item}" for item in fields.next_steps),
        }
        document = _open_document(self._template_bytes)
        seen: set[str] = set()
        for paragraph in _paragraphs(document):
            names = _PLACEHOLDER_RE.findall(paragraph.text)
            if not names:
                continue
            if len(names) != 1 or paragraph.text.strip() != "{{" + names[0] + "}}":
                raise DocumentTemplateError("template placeholders must occupy one paragraph")
            name = names[0]
            if name not in replacements or name in seen:
                raise DocumentTemplateError("template contains an unknown or repeated placeholder")
            paragraph.text = replacements[name]
            seen.add(name)
        if seen != _PROJECT_BRIEF_PLACEHOLDERS:
            raise DocumentTemplateError("template did not consume every approved field")

        output = io.BytesIO()
        document.save(output)
        content = output.getvalue()
        if not content or len(content) > _MAX_DOCX_BYTES:
            raise DocumentTemplateError("rendered DOCX exceeds the output limit")
        return self.verify(content, arguments)

    def verify(
        self,
        content: bytes,
        value: TemplateDocumentArguments | dict[str, Any],
    ) -> RenderedDocument:
        arguments = (
            value if isinstance(value, TemplateDocumentArguments)
            else self.validate_arguments(value)
        )
        descriptor = self.descriptor(arguments.template_id)
        if arguments.template_sha256 != descriptor.content_sha256:
            raise DocumentTemplateError("approved template digest does not match the server template")
        _validate_package(content, require_placeholders=False)
        reopened = _open_document(content)
        if _placeholders(reopened):
            raise DocumentTemplateError("rendered DOCX still contains template placeholders")
        text = _document_text(reopened)
        fields = arguments.fields
        expected = (
            fields.title,
            fields.prepared_for,
            fields.prepared_on,
            fields.summary,
            *fields.objectives,
            *fields.next_steps,
        )
        if any(item not in text for item in expected):
            raise DocumentTemplateError("rendered DOCX failed expected-text read-back")
        return RenderedDocument(
            content=content,
            content_sha256=hashlib.sha256(content).hexdigest(),
            text=text,
            expected_text=expected,
        )


def _plain_text(value: str, *, maximum: int) -> str:
    cleaned = value.strip()
    if not cleaned or len(cleaned) > maximum:
        raise ValueError(f"plain text must contain 1 to {maximum} characters")
    if any(ord(char) < 32 and char not in {"\n", "\t"} for char in cleaned):
        raise ValueError("plain text contains control characters")
    return cleaned


def _open_document(content: bytes) -> DocumentObject:
    try:
        return Document(io.BytesIO(content))
    except Exception as exc:
        raise DocumentTemplateError("DOCX package could not be reopened") from exc


def _paragraphs(container: DocumentObject | _Cell) -> list[Paragraph]:
    paragraphs = list(container.paragraphs)
    for table in container.tables:
        paragraphs.extend(_table_paragraphs(table))
    return paragraphs


def _table_paragraphs(table: Table) -> list[Paragraph]:
    paragraphs: list[Paragraph] = []
    for row in table.rows:
        for cell in row.cells:
            paragraphs.extend(_paragraphs(cell))
    return paragraphs


def _placeholders(document: DocumentObject) -> frozenset[str]:
    names: set[str] = set()
    for paragraph in _paragraphs(document):
        names.update(_PLACEHOLDER_RE.findall(paragraph.text))
    return frozenset(names)


def _document_text(document: DocumentObject) -> str:
    return "\n".join(
        paragraph.text.strip()
        for paragraph in _paragraphs(document)
        if paragraph.text.strip()
    )


def _validate_package(content: bytes, *, require_placeholders: bool) -> None:
    if len(content) > _MAX_DOCX_BYTES:
        raise DocumentTemplateError("DOCX package exceeds the output limit")
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            infos = archive.infolist()
            names = {info.filename for info in infos}
            if len(infos) > _MAX_PACKAGE_FILES:
                raise DocumentTemplateError("DOCX package contains too many parts")
            if sum(info.file_size for info in infos) > _MAX_PACKAGE_UNCOMPRESSED_BYTES:
                raise DocumentTemplateError("DOCX package expands beyond its limit")
            if "[Content_Types].xml" not in names or "word/document.xml" not in names:
                raise DocumentTemplateError("DOCX package is missing required parts")
            for name in names:
                path = Path(name)
                if path.is_absolute() or ".." in path.parts or "\\" in name:
                    raise DocumentTemplateError("DOCX package contains an unsafe path")
                lowered = name.casefold()
                if (
                    "vbaproject" in lowered
                    or "activex" in lowered
                    or lowered.startswith("customui/")
                    or lowered.endswith((
                        ".bmp", ".gif", ".jpeg", ".jpg", ".png", ".svg", ".tif",
                        ".tiff", ".webp",
                    ))
                    or lowered.startswith("word/media/")
                    or lowered.startswith("word/embeddings/")
                ):
                    raise DocumentTemplateError("DOCX package contains forbidden active content")
            content_types = archive.read("[Content_Types].xml").lower()
            if b"macroenabled" in content_types or b"activex" in content_types:
                raise DocumentTemplateError("DOCX package contains forbidden active content")
            for name in names:
                if not name.casefold().endswith(".rels"):
                    continue
                if _EXTERNAL_RELATIONSHIP_RE.search(archive.read(name)):
                    raise DocumentTemplateError("DOCX package contains an external relationship")
            for name in names:
                if not name.startswith("word/") or not name.endswith(".xml"):
                    continue
                xml = archive.read(name)
                if b"<w:instrText" in xml or b"<w:fldSimple" in xml:
                    raise DocumentTemplateError("DOCX package contains fields or formulas")
            if require_placeholders and b"{{title}}" not in archive.read("word/document.xml"):
                raise DocumentTemplateError("DOCX template does not contain its title placeholder")
    except zipfile.BadZipFile as exc:
        raise DocumentTemplateError("DOCX package is malformed") from exc


__all__ = [
    "DOCX_MIME",
    "PROJECT_BRIEF_TEMPLATE_ID",
    "TEMPLATE_OPERATION",
    "TEMPLATE_WORKER_VERSION",
    "DocumentTemplateCatalog",
    "DocumentTemplateError",
    "ProjectBriefFields",
    "RenderedDocument",
    "TemplateDescriptor",
    "TemplateDocumentArguments",
]

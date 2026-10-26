"""Bounded, in-memory document generation from declarative content.

These renderers never evaluate formulas, run supplied code, open paths, or fetch
remote resources. The caller owns authentication, persistence and download URLs.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import re
import unicodedata
from dataclasses import dataclass
from typing import Annotated, Any, Literal, Self

from docx import Document
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment
from pydantic import BaseModel, ConfigDict, Field, field_validator, model_validator

MAX_INPUT_BYTES = 256 * 1024
MAX_OUTPUT_BYTES = 4 * 1024 * 1024
MAX_ROWS = 1000
MAX_COLUMNS = 32
MAX_CELL_CHARS = 32767
_MAX_JSON_DEPTH = 24
_MAX_JSON_NODES = 32000
_INVALID_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")
_INVALID_FILENAME = re.compile(r'[<>:"/\\|?*]')
_INVALID_SHEET = re.compile(r"[\[\]:*?/\\]")
_RESERVED_NAMES = {"con", "prn", "aux", "nul"} | {
    f"{prefix}{number}" for prefix in ("com", "lpt") for number in range(1, 10)
}

Scalar = str | int | float | bool | None
Row = Annotated[list[Scalar], Field(max_length=MAX_COLUMNS)]
Rows = Annotated[list[Row], Field(min_length=1, max_length=MAX_ROWS)]
FileFormat = Literal["txt", "md", "csv", "json", "docx", "xlsx"]

_MIMES = {
    "txt": "text/plain",
    "md": "text/markdown",
    "csv": "text/csv",
    "json": "application/json",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


def _check_rows(rows: list[list[Scalar]]) -> None:
    for row in rows:
        for value in row:
            if isinstance(value, str) and (
                len(value) > MAX_CELL_CHARS or _INVALID_XML.search(value)
            ):
                raise ValueError("A cell contains unsupported text or exceeds the cell limit.")


def _check_json(value: Any) -> None:
    """Accept JSON values only; bound nesting and work before serialization."""
    pending = [(value, 0)]
    nodes = 0
    while pending:
        item, depth = pending.pop()
        nodes += 1
        if depth > _MAX_JSON_DEPTH or nodes > _MAX_JSON_NODES:
            raise ValueError("JSON content exceeds the nesting or item limit.")
        if item is None or isinstance(item, (str, bool, int)):
            continue
        if isinstance(item, float):
            if not math.isfinite(item):
                raise ValueError("JSON numbers must be finite.")
        elif isinstance(item, list):
            pending.extend((child, depth + 1) for child in item)
        elif isinstance(item, dict) and all(isinstance(key, str) for key in item):
            pending.extend((child, depth + 1) for child in item.values())
        else:
            raise ValueError("JSON content must contain only JSON values.")


class FileSheet(BaseModel):
    """One worksheet of literal scalar cells; formulas are never evaluated."""

    model_config = ConfigDict(extra="forbid", strict=True, allow_inf_nan=False,
                              hide_input_in_errors=True)
    name: str = Field(min_length=1, max_length=31)
    rows: Rows

    @field_validator("name")
    @classmethod
    def check_name(cls, value: str) -> str:
        if (not value.strip() or value != value.strip() or _INVALID_SHEET.search(value)
                or value.startswith("'") or value.endswith("'")
                or any(unicodedata.category(char).startswith("C") for char in value)):
            raise ValueError("Use a nonempty worksheet name without reserved characters.")
        return value

    @field_validator("rows")
    @classmethod
    def check_rows(cls, value: list[list[Scalar]]) -> list[list[Scalar]]:
        _check_rows(value)
        return value


class CreateFileRequest(BaseModel):
    """Supply exactly the content field belonging to the chosen format."""

    model_config = ConfigDict(extra="forbid", strict=True, allow_inf_nan=False,
                              hide_input_in_errors=True)
    filename: str = Field(min_length=1, max_length=180, description="Display basename with format suffix.")
    format: FileFormat
    content: str | None = Field(default=None, max_length=200000,
                               description="Plain text for txt, md or docx.")
    rows: Rows | None = Field(default=None, description="CSV rows, including headers if wanted.")
    data: Any = Field(default=None, description="Any JSON value; used only with json format.")
    sheets: list[FileSheet] | None = Field(default=None, min_length=1, max_length=8,
                                          description="Named XLSX sheets with literal cells.")

    @field_validator("filename")
    @classmethod
    def check_filename(cls, value: str) -> str:
        if (value != value.strip() or value.startswith(".") or value.endswith(".")
                or _INVALID_FILENAME.search(value)
                or any(unicodedata.category(char).startswith("C") for char in value)
                or value.split(".", 1)[0].casefold() in _RESERVED_NAMES):
            raise ValueError("Use a safe filename without directories or reserved characters.")
        return value

    @field_validator("content")
    @classmethod
    def check_content(cls, value: str | None) -> str | None:
        if value is not None and _INVALID_XML.search(value):
            raise ValueError("Content contains unsupported control characters.")
        return value

    @field_validator("rows")
    @classmethod
    def check_rows(cls, value: list[list[Scalar]] | None) -> list[list[Scalar]] | None:
        if value is not None:
            _check_rows(value)
        return value

    @model_validator(mode="after")
    def check_format_and_size(self) -> Self:
        field = {"txt": "content", "md": "content", "docx": "content", "csv": "rows",
                 "json": "data", "xlsx": "sheets"}[self.format]
        supplied = self.model_fields_set & {"content", "rows", "data", "sheets"}
        if supplied != {field} or (field != "data" and getattr(self, field) is None):
            raise ValueError("Supply only the content field required by this format.")
        suffix = "." + self.format
        if not self.filename.casefold().endswith(suffix) or not self.filename[:-len(suffix)].strip(" ."):
            raise ValueError("The filename must have the selected format's suffix.")
        if self.format == "json":
            _check_json(self.data)
        if self.sheets is not None:
            names = [sheet.name.casefold() for sheet in self.sheets]
            if len(set(names)) != len(names):
                raise ValueError("Worksheet names must be unique.")
            if sum(len(sheet.rows) for sheet in self.sheets) > MAX_ROWS:
                raise ValueError("The workbook exceeds the total row limit.")
        try:
            size = len(json.dumps(self.model_dump(exclude_unset=True), ensure_ascii=False,
                                  allow_nan=False, separators=(",", ":")).encode("utf-8"))
        except (ValueError, TypeError, UnicodeError) as error:
            raise ValueError("Content cannot be represented as bounded UTF-8 JSON.") from error
        if size > MAX_INPUT_BYTES:
            raise ValueError("File content exceeds the total input limit.")
        return self


@dataclass(frozen=True)
class RenderedFile:
    data: bytes
    mime_type: str
    text: str
    sha256: str


def _csv_text(rows: list[list[Scalar]], *, protect_formulas: bool) -> str:
    target = io.StringIO(newline="")
    writer = csv.writer(target, quoting=csv.QUOTE_ALL, lineterminator="\n")
    for row in rows:
        cells = []
        for value in row:
            if isinstance(value, str) and protect_formulas and (
                value.lstrip().startswith(("=", "+", "-", "@"))
                or value.startswith(("\t", "\r", "\n"))
            ):
                value = "'" + value
            cells.append(value)
        writer.writerow(cells)
    return target.getvalue()


def _workbook(sheets: list[FileSheet]) -> bytes:
    workbook = Workbook(write_only=True)
    target = io.BytesIO()
    for spec in sheets:
        sheet = workbook.create_sheet(spec.name)
        for row in spec.rows:
            cells = []
            for value in row:
                # Excel numeric cells have only 15 significant digits. Preserve
                # larger integer values exactly rather than silently rounding.
                if isinstance(value, int) and not isinstance(value, bool) and abs(value) > 999999999999999:
                    value = str(value)
                cell = WriteOnlyCell(sheet, value=value)
                if isinstance(value, str):
                    cell.data_type = "s"
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                cells.append(cell)
            sheet.append(cells)
    workbook.save(target)
    workbook.close()
    return target.getvalue()


def render_file(request: CreateFileRequest) -> RenderedFile:
    """Render validated data entirely in memory; raise safe errors on limits."""
    # Revalidate at the boundary so model_construct/model_copy cannot bypass
    # the bounds if a future caller hands us an already-created model.
    request = CreateFileRequest.model_validate(request.model_dump(exclude_unset=True))
    match request.format:
        case "txt" | "md":
            text = request.content or ""
            data = text.encode("utf-8")
        case "json":
            text = json.dumps(request.data, ensure_ascii=False, allow_nan=False, indent=2) + "\n"
            data = text.encode("utf-8")
        case "csv":
            text = _csv_text(request.rows or [], protect_formulas=False)
            data = _csv_text(request.rows or [], protect_formulas=True).encode("utf-8-sig")
        case "docx":
            text = request.content or ""
            document = Document()
            for paragraph in text.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
                document.add_paragraph(paragraph)
            target = io.BytesIO()
            document.save(target)
            data = target.getvalue()
        case "xlsx":
            sheets = request.sheets or []
            text = "\n".join(f"{sheet.name}\n{_csv_text(sheet.rows, protect_formulas=False)}"
                             for sheet in sheets)
            data = _workbook(sheets)
    if len(data) > MAX_OUTPUT_BYTES or len(text.encode("utf-8")) > MAX_OUTPUT_BYTES:
        raise ValueError("The generated file exceeds the output limit.")
    return RenderedFile(data=data, mime_type=_MIMES[request.format], text=text,
                        sha256=hashlib.sha256(data).hexdigest())


__all__ = ["CreateFileRequest", "FileSheet", "RenderedFile", "render_file"]

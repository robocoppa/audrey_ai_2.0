"""Local, revision-specific XLSX projections of saved Private tax cases.

The canonical case remains authoritative. Snapshots have finite retention and
never enter upload ingestion. Authentication and a fresh canonical case check
belong to the caller, including a check after a download finishes rendering.
"""
from __future__ import annotations

import asyncio
import hashlib
import json
import os
import re
import shutil
import time
import uuid
import weakref
from dataclasses import dataclass, field
from datetime import UTC, datetime
from pathlib import Path
from typing import Any, BinaryIO

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from audrey.app_state.tax_cases import tax_review_scope

EXPORTER_VERSION = 1
_EXPORT_ID = re.compile(r"txlsx_[0-9a-f]{64}\Z")
_INVALID_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_NOTICE = (
    "Private server-side export. Selected sources and illustrative arithmetic only; "
    "not a completed return, legal accuracy check, or determination of eligibility. "
    "Check the sources and figures before filing. Spreadsheet edits do not update the saved case."
)


class TaxExportError(ValueError):
    """A public safe export failure; never include case contents in this error."""

    def __init__(self, code: str, status_code: int, detail: str):
        super().__init__(detail)
        self.code = code
        self.status_code = status_code
        self.detail = detail


@dataclass(frozen=True)
class TaxExportArtifact:
    path: Path
    filename: str
    case_revision: int
    review_revision: int | None
    file: BinaryIO
    size: int

    def close(self) -> None:
        self.file.close()


@dataclass
class _OwnerState:
    lock: asyncio.Lock = field(default_factory=asyncio.Lock)
    epoch: int = 0


def _not_found() -> TaxExportError:
    return TaxExportError("tax_export_not_found", 404, "Tax review file not found.")


def _unavailable() -> TaxExportError:
    return TaxExportError("tax_export_unavailable", 503, "The saved review is available, but its file could not be prepared. Try again.")


def _limit() -> TaxExportError:
    return TaxExportError("tax_export_limit", 422, "This review exceeds the file export limit. The saved review is unchanged.")


def _expired() -> TaxExportError:
    return TaxExportError("tax_export_expired", 410, "This review file has expired. Open the current saved review to request a new file.")


def _hash(value: str) -> str:
    return hashlib.sha256(value.encode("utf-8")).hexdigest()


def _encoded(value: Any) -> bytes:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False).encode("utf-8")


def _text(value: Any) -> str:
    if value is None:
        return "Unknown"
    if isinstance(value, str):
        return value
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    return str(value)


def _identity(case: dict) -> tuple[str, str, int]:
    try:
        case_id, conversation_id, revision = case["id"], case["conversation_id"], case["revision"]
        if (not isinstance(case_id, str) or not 0 < len(case_id) <= 200
                or not isinstance(conversation_id, str) or not 0 < len(conversation_id) <= 200
                or type(revision) is not int or revision < 1
                or type(case.get("tax_year")) is not int or not 1900 <= case["tax_year"] <= 2100):
            raise ValueError
    except (KeyError, TypeError, ValueError):
        raise _not_found() from None
    return case_id, conversation_id, revision


class TaxReviewExports:
    """Bounded disk cache with owner serialization and lazy off-loop rendering.

    One server process owns this cache, matching the application deployment.
    Cleanup invalidates queued work before waiting for its owner lock; a renderer
    cannot publish after deletion starts. Durable tombstones stop later reuse of
    a deleted case. Callers may undo a tombstone when canonical deletion fails.
    """

    def __init__(
        self, root: Path, *, retention_seconds: float = 30 * 24 * 60 * 60,
        max_case_exports: int = 8, max_snapshots: int = 512,
        max_cache_bytes: int = 128 * 1024 * 1024,
        max_snapshot_bytes: int = 512 * 1024, max_workbook_bytes: int = 8 * 1024 * 1024,
        max_rows: int = 4096, concurrency: int = 2,
    ):
        if min(retention_seconds, max_case_exports, max_snapshots, max_cache_bytes,
               max_snapshot_bytes, max_workbook_bytes, max_rows, concurrency) <= 0:
            raise ValueError("Export limits must be positive.")
        self.root = Path(root).resolve()
        self.retention_seconds = retention_seconds
        self.max_case_exports = max_case_exports
        self.max_snapshots = max_snapshots
        self.max_cache_bytes = max_cache_bytes
        self.max_snapshot_bytes = max_snapshot_bytes
        self.max_workbook_bytes = max_workbook_bytes
        self.max_rows = max_rows
        self._renders = asyncio.Semaphore(concurrency)
        self._storage = asyncio.Lock()
        self._owners: weakref.WeakValueDictionary[str, _OwnerState] = weakref.WeakValueDictionary()
        self._active_temporary: set[Path] = set()

    def _owner_state(self, user_id: str) -> _OwnerState:
        key = _hash(user_id)
        state = self._owners.get(key)
        if state is None:
            state = _OwnerState()
            self._owners[key] = state
        return state

    def _paths(self, user_id: str, case_id: str) -> tuple[Path, Path]:
        owner = self.root / _hash(user_id)
        return owner, owner / _hash(case_id)

    @staticmethod
    def _check_epoch(state: _OwnerState, epoch: int) -> None:
        if state.epoch != epoch:
            raise _not_found()

    async def _disk(self, function, *args):
        async with self._storage:
            task = asyncio.create_task(asyncio.to_thread(function, *args))
            try:
                return await asyncio.shield(task)
            except asyncio.CancelledError:
                # A filesystem worker cannot be cancelled. Keep the storage and
                # owner locks until it exits, so lifecycle cleanup runs last.
                result = await task
                if isinstance(result, TaxExportArtifact):
                    result.close()
                raise

    async def prepare(self, *, user_id: str, case: dict) -> dict:
        """Capture a bounded immutable revision; no workbook or model calls yet."""
        if not user_id:
            raise _not_found()
        case_id, _, revision = _identity(case)
        if not isinstance(case.get("review"), dict) or not case["review"]:
            raise _not_found()
        state = self._owner_state(user_id)
        epoch = state.epoch
        # Encode before yielding so another operation cannot mutate this reference
        # between the saved revision and its retained snapshot.
        try:
            encoded_case = _encoded(case)
        except (TypeError, ValueError, RecursionError):
            raise _limit() from None
        if len(encoded_case) > self.max_snapshot_bytes:
            raise _limit()
        export_id = "txlsx_" + _hash(f"{EXPORTER_VERSION}\0{user_id}\0{case_id}\0{revision}")
        async with state.lock:
            self._check_epoch(state, epoch)
            try:
                result = await self._disk(self._prepare_sync, user_id, encoded_case, export_id)
                self._check_epoch(state, epoch)
                return result
            except TaxExportError:
                raise
            except (OSError, ValueError, KeyError, TypeError):
                raise _unavailable() from None

    def _prepare_sync(self, user_id: str, encoded_case: bytes, export_id: str) -> dict:
        case = json.loads(encoded_case)
        case_id, conversation_id, revision = _identity(case)
        owner, directory = self._paths(user_id, case_id)
        self._not_deleted(owner, directory)
        self._mkdir(directory)
        path = directory / f"{export_id}.json"
        now = time.time()
        self._prune(now, protected=path)
        if path.exists():
            snapshot = self._load(path, user_id, case)
            if snapshot["expires_at"] > now:
                return self._reference(snapshot)
            self._remove_pair(path)
        review_revision = case["review"].get("review_revision")
        filename = f"tax-review-{case['tax_year']}-case-r{revision}-review-r{review_revision or 'unknown'}.xlsx"
        self._mkdir(directory)
        snapshot = {
            "schema": EXPORTER_VERSION, "id": export_id, "owner": user_id,
            "case_id": case_id, "conversation_id": conversation_id,
            "case_revision": revision, "review_revision": review_revision,
            "created_at": now, "expires_at": now + self.retention_seconds,
            "filename": filename, "case": case,
        }
        payload = _encoded(snapshot)
        if len(payload) > self.max_snapshot_bytes:
            raise _limit()
        self._reserve(len(payload), path)
        self._atomic(path, payload)
        self._prune(now, protected=path)
        return self._reference(snapshot)

    @staticmethod
    def _reference(snapshot: dict) -> dict:
        return {
            "id": snapshot["id"], "case_revision": snapshot["case_revision"],
            "review_revision": snapshot["review_revision"], "filename": snapshot["filename"],
            "expires_at": datetime.fromtimestamp(snapshot["expires_at"], UTC).isoformat(),
            "download_url": f"/api/conversations/{snapshot['conversation_id']}/tax-review/exports/{snapshot['id']}/download",
        }

    async def download(self, *, user_id: str, case: dict, export_id: str) -> TaxExportArtifact:
        """Render exactly the retained snapshot, allowing a newer current case."""
        case_id, _, _ = _identity(case)
        if not user_id or not isinstance(export_id, str) or not _EXPORT_ID.fullmatch(export_id):
            raise _not_found()
        state = self._owner_state(user_id)
        epoch = state.epoch
        _, directory = self._paths(user_id, case_id)
        snapshot_path = directory / f"{export_id}.json"
        temporary = directory / f".{export_id}-{uuid.uuid4().hex}.tmp"
        async with state.lock:
            self._check_epoch(state, epoch)
            try:
                snapshot, cached = await self._disk(self._download_sync, snapshot_path, user_id, case)
                if not cached:
                    async with self._renders:
                        self._check_epoch(state, epoch)
                        self._active_temporary.add(temporary)
                        task = asyncio.create_task(asyncio.to_thread(self._render, snapshot, temporary))
                        try:
                            await asyncio.shield(task)
                        except asyncio.CancelledError:
                            # to_thread cannot be stopped; wait before releasing
                            # the lifecycle lock so cleanup wins over the writer.
                            try:
                                await task
                            finally:
                                await asyncio.to_thread(temporary.unlink, missing_ok=True)
                            raise
                    self._check_epoch(state, epoch)
                    await self._disk(self._publish, snapshot_path, temporary, user_id, case)
                self._check_epoch(state, epoch)
                artifact = await self._disk(self._open_artifact, snapshot_path, user_id, case)
                try:
                    self._check_epoch(state, epoch)
                except TaxExportError:
                    artifact.close()
                    raise
                return artifact
            except TaxExportError:
                raise
            except Exception:  # noqa: BLE001 - optional renderer; hide content-bearing library errors.
                raise _unavailable() from None
            finally:
                await asyncio.to_thread(temporary.unlink, missing_ok=True)
                self._active_temporary.discard(temporary)

    def _open_artifact(self, path: Path, user_id: str, case: dict) -> TaxExportArtifact:
        snapshot = self._load(path, user_id, case)
        if snapshot["expires_at"] <= time.time():
            self._remove_pair(path)
            raise _expired()
        output = path.with_suffix(".xlsx")
        if output.is_symlink():
            raise _not_found()
        handle = output.open("rb")
        try:
            size = os.fstat(handle.fileno()).st_size
            if not 0 < size <= self.max_workbook_bytes:
                raise _limit()
            return TaxExportArtifact(
                path=output, filename=snapshot["filename"],
                case_revision=snapshot["case_revision"], review_revision=snapshot["review_revision"],
                file=handle, size=size,
            )
        except BaseException:
            handle.close()
            raise

    def _download_sync(self, path: Path, user_id: str, case: dict) -> tuple[dict, bool]:
        snapshot = self._load(path, user_id, case)
        if snapshot["expires_at"] <= time.time():
            self._remove_pair(path)
            raise _expired()
        output = path.with_suffix(".xlsx")
        if output.is_symlink():
            raise _not_found()
        if output.exists() and not 0 < output.stat().st_size <= self.max_workbook_bytes:
            output.unlink()
        return snapshot, output.exists()

    def _load(self, path: Path, user_id: str, case: dict) -> dict:
        case_id, conversation_id, _ = _identity(case)
        owner, directory = self._paths(user_id, case_id)
        self._not_deleted(owner, directory)
        if not path.is_file() or path.is_symlink():
            raise _not_found()
        if path.stat().st_size > self.max_snapshot_bytes:
            raise _limit()
        snapshot = json.loads(path.read_bytes())
        if (snapshot.get("schema") != EXPORTER_VERSION or snapshot.get("owner") != user_id
                or snapshot.get("case_id") != case_id or snapshot.get("conversation_id") != conversation_id
                or snapshot.get("id") != path.stem or _identity(snapshot["case"]) != (
                    case_id, conversation_id, snapshot.get("case_revision"))):
            raise _not_found()
        return snapshot

    def _publish(self, snapshot_path: Path, temporary: Path, user_id: str, case: dict) -> None:
        snapshot = self._load(snapshot_path, user_id, case)
        if snapshot["expires_at"] <= time.time():
            self._remove_pair(snapshot_path)
            raise _expired()
        size = temporary.stat().st_size
        if not 0 < size <= self.max_workbook_bytes:
            raise _limit()
        self._reserve(size, snapshot_path)
        temporary.replace(snapshot_path.with_suffix(".xlsx"))

    @staticmethod
    def _mkdir(path: Path) -> None:
        path.mkdir(mode=0o700, parents=True, exist_ok=True)
        # parents=True applies its default mode to parents, so set every cache
        # directory explicitly before private bytes are written.
        path.chmod(0o700)
        path.parent.chmod(0o700)
        path.parent.parent.chmod(0o700)

    @staticmethod
    def _not_deleted(owner: Path, directory: Path) -> None:
        if owner.is_symlink() or directory.is_symlink() or (owner / ".deleted").exists() or (directory / ".deleted").exists():
            raise _not_found()

    @staticmethod
    def _atomic(path: Path, payload: bytes) -> None:
        temporary = path.with_name(f".{path.name}-{uuid.uuid4().hex}.tmp")
        try:
            with temporary.open("xb") as output:
                temporary.chmod(0o600)
                output.write(payload)
                output.flush()
                os.fsync(output.fileno())
            temporary.replace(path)
        finally:
            temporary.unlink(missing_ok=True)

    @staticmethod
    def _empty_directories(directory: Path) -> None:
        for path in (directory, directory.parent):
            try:
                path.rmdir()
            except OSError:
                pass  # Another retained snapshot, marker or renderer still owns it.

    @staticmethod
    def _remove_pair(path: Path) -> None:
        path.unlink(missing_ok=True)
        path.with_suffix(".xlsx").unlink(missing_ok=True)
        TaxReviewExports._empty_directories(path.parent)

    def _inventory(self) -> list[tuple[float, Path, int]]:
        result = []
        if self.root.exists():
            for path in self.root.glob("*/*/txlsx_*.json"):
                if not path.is_file() or path.is_symlink():
                    continue
                size = path.stat().st_size
                output = path.with_suffix(".xlsx")
                if output.is_file():
                    size += output.stat().st_size
                result.append((path.stat().st_mtime, path, size))
        return sorted(result)

    def _prune_markers(self, now: float) -> None:
        markers = sorted((path.stat().st_mtime, path)
                         for path in self.root.glob("*/*/.deleted") if path.is_file())
        for index, (modified, marker) in enumerate(markers):
            if now - modified >= self.retention_seconds or index < len(markers) - self.max_snapshots:
                marker.unlink(missing_ok=True)
                self._empty_directories(marker.parent)

    def _prune(self, now: float, *, protected: Path) -> None:
        self._prune_markers(now)
        # A killed process may leave an unpublished output. There is only one
        # cache-owning application process; live render paths are retained.
        for temporary in self.root.glob("*/*/.*.tmp"):
            if temporary not in self._active_temporary:
                temporary.unlink(missing_ok=True)
                self._empty_directories(temporary.parent)
        entries = self._inventory()
        per_case: dict[Path, int] = {protected.parent: 1} if protected.exists() else {}
        for modified, path, _ in reversed(entries):
            if path == protected:
                continue
            per_case[path.parent] = per_case.get(path.parent, 0) + 1
            if (
                now - modified >= self.retention_seconds
                or per_case[path.parent] > self.max_case_exports
            ):
                self._remove_pair(path)
        entries = self._inventory()
        while len(entries) > self.max_snapshots:
            _, path, _ = next(entry for entry in entries if entry[1] != protected)
            self._remove_pair(path)
            entries = [entry for entry in entries if entry[1] != path]

    def _reserve(self, added_bytes: int, protected: Path) -> None:
        entries = self._inventory()
        size = sum(item[2] for item in entries)
        for _, path, pair_bytes in entries:
            if size + added_bytes <= self.max_cache_bytes:
                return
            if path != protected:
                self._remove_pair(path)
                size -= pair_bytes
        if size + added_bytes > self.max_cache_bytes:
            raise _limit()

    async def delete_case(self, *, user_id: str, case_id: str) -> None:
        await self._cleanup(user_id, case_id, restore=False)

    async def allow_case(self, *, user_id: str, case_id: str) -> None:
        """Roll back the derived-data tombstone after canonical deletion fails."""
        await self._cleanup(user_id, case_id, restore=True)

    async def delete_owner(self, *, user_id: str) -> None:
        await self._cleanup(user_id, None, restore=False)

    async def allow_owner(self, *, user_id: str) -> None:
        await self._cleanup(user_id, None, restore=True)

    async def _cleanup(self, user_id: str, case_id: str | None, *, restore: bool) -> None:
        state = self._owner_state(user_id)
        state.epoch += 1
        async with state.lock:
            try:
                await self._disk(self._cleanup_sync, user_id, case_id, restore)
            except OSError:
                raise _unavailable() from None

    def _cleanup_sync(self, user_id: str, case_id: str | None, restore: bool) -> None:
        owner, directory = self._paths(user_id, case_id or "")
        target = directory if case_id is not None else owner
        if target.is_symlink():
            raise _not_found()
        if restore:
            (target / ".deleted").unlink(missing_ok=True)
            return
        retained = [item.name for item in owner.iterdir()
                    if item.is_dir() and not item.is_symlink()
                    and re.fullmatch(r"[0-9a-f]{64}", item.name)] if case_id is None and owner.exists() else []
        if target.exists():
            shutil.rmtree(target)
        if case_id is not None:
            self._mkdir(directory)
            (directory / ".deleted").touch(mode=0o600)
        else:
            # Personal-data purges preserve the account. Block its old case IDs,
            # while a genuinely new case can later prepare an export normally.
            for name in retained:
                tombstone = owner / name
                self._mkdir(tombstone)
                (tombstone / ".deleted").touch(mode=0o600)
        self._prune_markers(time.time())

    def _render(self, snapshot: dict, destination: Path) -> None:
        """Only code renders recorded values; no formulas, macros or URL fetches."""
        workbook = Workbook(write_only=True)
        rows_written = 0

        def sheet(name: str, widths: list[int]):
            result = workbook.create_sheet(name)
            result.freeze_panes = "A2"
            for index, width in enumerate(widths, 1):
                result.column_dimensions[get_column_letter(index)].width = width
            return result

        def append(target, values, *, header=False):
            nonlocal rows_written
            rows_written += 1
            if rows_written > self.max_rows:
                raise _limit()
            cells = []
            for value in values:
                value = _text(value)
                if len(value) > 32767 or _INVALID_XML.search(value):
                    raise _limit()
                cell = WriteOnlyCell(target, value=value)
                # Assign the type after value: openpyxl otherwise turns '=' into
                # a formula. Decimal strings also stay exact beyond Excel's
                # numeric precision, including trailing zeros and large values.
                cell.data_type = "s"
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if header:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="203D60")
                cells.append(cell)
            target.append(cells)

        try:
            case = snapshot["case"]
            review = case["review"]
            overview = sheet("Overview", [30, 100])
            append(overview, ["Tax review", "Saved case snapshot"], header=True)
            for key, value in (
                ("Tax year", case.get("tax_year")), ("Jurisdiction", case.get("jurisdiction")),
                ("Goal", case.get("goal")), ("Case revision", snapshot["case_revision"]),
                ("Review revision", snapshot["review_revision"]), ("Review saved", review.get("saved_at")),
                ("Reference release", case.get("reference_release")), ("Summary", review.get("summary")),
                ("Review status", review.get("status", "Saved; see unresolved items")),
                ("Review stale", review.get("stale", review.get("review_revision") != case.get("revision"))),
                ("Scope", tax_review_scope()["notice"]), ("Notice", _NOTICE),
                ("File expires", datetime.fromtimestamp(snapshot["expires_at"], UTC).isoformat()),
                ("Next question", review.get("next_question") or "None recorded"),
            ):
                append(overview, [key, value])
            append(overview, ["Finding", "Recorded detail and supporting IDs"], header=True)
            for finding in review.get("findings", []):
                append(overview, [finding.get("label"), finding.get("detail")])
                append(overview, ["Supporting records", {key: finding.get(key, []) for key in (
                    "fact_ids", "evidence_ids", "calculation_ids", "return_item_ids") }])
            append(overview, ["Proposed return item", "Status and supporting IDs"], header=True)
            for item in case.get("return_items", []):
                append(overview, [item.get("label"), item])
            facts = sheet("Facts", [24, 40, 70, 20, 18, 18, 80, 80])
            append(facts, ["ID", "Fact", "Saved value", "Status", "Original year", "Confirmed year", "Provenance", "Correction history"], header=True)
            for fact in case.get("facts", []):
                append(facts, [fact.get("id"), fact.get("label"), fact.get("value"), fact.get("status"),
                    fact.get("case_tax_year"), fact.get("confirmed_tax_year"), fact.get("provenance", fact.get("user_statement")), fact.get("history", [])])
            evidence = sheet("Evidence", [24, 40, 45, 18, 24, 24, 18, 60, 110, 80, 35, 22])
            append(evidence, ["ID", "Claim", "Official document", "Page / section", "Revision", "Stated tax year", "Year status", "Original citation", "Retained passage", "Document ID / release", "Conflicts / superseded by", "Status"], header=True)
            sources = case.get("export_sources", {})
            for item in case.get("evidence", []):
                source = sources.get(item.get("id"), sources.get(item.get("document_id"), {}))
                citation = source.get("source_url") or item.get("source_url") or item.get("url")
                append(evidence, [item.get("id"), item.get("claim"),
                    source.get("title") or source.get("product_number") or item.get("document_id"),
                    f"PDF page {item.get('page', 'Unknown')} / {item.get('section') or 'No section recorded'}",
                    source.get("revision_date") or source.get("revision") or item.get("revision"), item.get("tax_year"), item.get("tax_year_status"),
                    citation or "No original URL retained", item.get("passage"),
                    {"document_id": item.get("document_id"), "revision_sha256": item.get("revision"), "reference_release": item.get("reference_release")},
                    {"conflicts_with": item.get("conflicts_with", []), "superseded_by": item.get("superseded_by")}, item.get("status")])
            calculations = sheet("Calculations", [24, 40, 80, 55, 30, 50, 80, 22, 18])
            append(calculations, ["ID", "Calculation", "Named decimal inputs", "Expression (text)", "Rounding", "Recorded result (exact text)", "Supporting IDs", "Status", "Original year"], header=True)
            for item in case.get("calculations", []):
                append(calculations, [item.get("id"), item.get("label"), item.get("inputs"), item.get("expression"),
                    item.get("rounding"), item.get("result"), {"fact_ids": item.get("fact_ids", []), "evidence_ids": item.get("evidence_ids", []), "superseded_by": item.get("superseded_by")}, item.get("status"), item.get("case_tax_year")])
            unresolved = sheet("Unresolved items", [26, 55, 100, 28])
            append(unresolved, ["Kind / ID", "Topic", "Remaining question or limitation", "Status"], header=True)
            for item in case.get("missing_items", []):
                append(unresolved, [item.get("id"), item.get("label"), item.get("question"), item.get("status")])
            for index, limitation in enumerate(review.get("limitations", []), 1):
                append(unresolved, [f"Review limitation {index}", "Review limitation", limitation, "Unresolved"])
            for category in ("facts", "evidence", "calculations", "return_items"):
                for item in case.get(category, []):
                    status = item.get("status")
                    conflict = item.get("conflicts_with")
                    if status in {"proposed", "needs_recheck", "rejected"} or conflict:
                        append(unresolved, [item.get("id"), item.get("label") or item.get("claim"),
                            {"category": category, "conflicts_with": conflict or [], "recheck_for_tax_year": item.get("recheck_for_tax_year"), "recheck_for_reference_release": item.get("recheck_for_reference_release")}, status])
            with destination.open("xb") as output:
                destination.chmod(0o600)
                workbook.save(output)
                output.flush()
                os.fsync(output.fileno())
            if destination.stat().st_size > self.max_workbook_bytes:
                raise _limit()
        finally:
            # Close write-only worksheet generators even when a size/cell limit
            # interrupts rendering, including their temporary XML files.
            for worksheet in workbook.worksheets:
                if not worksheet.closed:
                    worksheet.close()
                writer = getattr(worksheet, "_writer", None)
                if writer is not None and Path(writer.out).exists():
                    writer.cleanup()
            workbook.close()

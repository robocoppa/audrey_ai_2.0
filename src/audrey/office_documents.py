"""Conservative Markdown-to-Word formatting; unsupported syntax stays literal.

Only text, native paragraphs and tables are emitted. No HTML, hyperlinks,
images, templates, supplied XML or external relationships are interpreted.
"""
from __future__ import annotations

import re

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.table import Table

_INLINE = re.compile(r'(\*\*[^*\n]+\*\*|\*[^*\n]+\*|`[^`\n]+`)')
_HEADING = re.compile(r'^(#{1,3})\s+(.+)$')
_LIST = re.compile(r'^\s{0,3}(?:(?P<bullet>[-+*])\s+|(?P<number>\d{1,4})[.)]\s+)(?P<text>.+)$')
_SEPARATOR = re.compile(r'^:?-{3,}:?$')
_MAX_BLOCKS = 2000
_MAX_TABLE_ROWS = 1000
_MAX_TABLE_COLUMNS = 32


def _inline(paragraph, text: str) -> None:
    position = 0
    for match in _INLINE.finditer(text):
        paragraph.add_run(text[position:match.start()])
        value = match.group()
        if value.startswith('**'):
            paragraph.add_run(value[2:-2]).bold = True
        elif value.startswith('*'):
            paragraph.add_run(value[1:-1]).italic = True
        else:
            paragraph.add_run(value[1:-1], style='InlineCode')
        position = match.end()
    paragraph.add_run(text[position:])


def _number(document, paragraph, start: int) -> None:
    # Use a local numbering instance so a source list starting at 4 stays at 4,
    # and separate lists do not silently continue an earlier list in Word.
    numbering = document.part.numbering_part.element
    base = document.styles['List Number'].element.pPr.numPr.numId.val
    definition = numbering.num_having_numId(base).abstractNumId.val
    instance = numbering.add_num(definition)
    instance.add_lvlOverride(ilvl=0).add_startOverride(start)
    properties = paragraph._p.get_or_add_pPr().get_or_add_numPr()
    properties.get_or_add_numId().val = instance.numId
    properties.get_or_add_ilvl().val = 0


def _cells(line: str) -> list[str] | None:
    if '|' not in line:
        return None
    value = line.strip().removeprefix('|').removesuffix('|')
    # Escaped pipes stay inside their own cell; do not support nested tables.
    return [part.strip().replace(r'\|', '|') for part in re.split(r'(?<!\\)\|', value)]


def _configure(document) -> None:
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.8)
    section.left_margin = section.right_margin = Inches(0.8)
    normal = document.styles['Normal']
    normal.font.name, normal.font.size = 'Aptos', Pt(11)
    normal.font.color.rgb = RGBColor.from_string('25364D')
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15
    for level, size in enumerate((26, 20, 15), 1):
        heading = document.styles[f'Heading {level}']
        heading.font.name, heading.font.size = 'Aptos', Pt(size)
        heading.font.color.rgb = RGBColor.from_string('16324F')
        heading.paragraph_format.keep_with_next = True
    for name, kind in [('CodeBlock', WD_STYLE_TYPE.PARAGRAPH), ('InlineCode', WD_STYLE_TYPE.CHARACTER)]:
        style = document.styles.add_style(name, kind)
        style.font.name, style.font.size = 'Consolas', Pt(10)
    document.styles['CodeBlock'].paragraph_format.space_after = Pt(8)


def _table(document, rows: list[list[str]]) -> None:
    width = len(rows[0])
    table = document.add_table(rows=0, cols=width)
    table.style = 'Table Grid'
    for index, values in enumerate(rows):
        cells = table.add_row().cells
        for cell, value in zip(cells, values, strict=True):
            paragraph = cell.paragraphs[0]
            _inline(paragraph, value)
            if index == 0:
                for run in paragraph.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
                shading = OxmlElement('w:shd')
                shading.set(qn('w:fill'), '16324F')
                cell._tc.get_or_add_tcPr().append(shading)
        if index == 0:
            repeat = OxmlElement('w:tblHeader')
            table.rows[0]._tr.get_or_add_trPr().append(repeat)


def _plain_text(document) -> str:
    blocks = []
    for block in document.iter_inner_content():
        if isinstance(block, Table):
            blocks.append('\n'.join('\t'.join(cell.text for cell in row.cells) for row in block.rows))
        else:
            blocks.append(block.text)
    return '\n\n'.join(blocks)


def markdown_document(content: str):
    """Return a native Word document and a plain projection of its actual body.

    This deliberately supports a small predictable subset rather than claiming
    arbitrary Markdown/HTML fidelity. Source links and image syntax stay text.
    """
    document = Document()
    _configure(document)
    lines = content.replace('\r\n', '\n').replace('\r', '\n').split('\n')
    index, count = 0, 0
    while index < len(lines):
        line = lines[index]
        if not line.strip():
            index += 1
            continue
        count += 1
        if count > _MAX_BLOCKS:
            raise ValueError('The document exceeds the formatting block limit.')
        fence = re.match(r'^\s{0,3}(`{3,}|~{3,})[^`~]*$', line)
        if fence:
            delimiter, code = fence.group(1), []
            index += 1
            while index < len(lines) and not re.match(r'^\s{0,3}' + re.escape(delimiter) + r'\s*$', lines[index]):
                code.append(lines[index])
                index += 1
            if index < len(lines):
                index += 1
                document.add_paragraph('\n'.join(code), style='CodeBlock')
            else:
                # Unclosed code fences are literal, including the opener.
                document.add_paragraph('\n'.join([line, *code]))
            continue
        heading = _HEADING.match(line)
        if heading:
            _inline(document.add_paragraph(style=f'Heading {len(heading.group(1))}'),
                    re.sub(r'\s+#+\s*$', '', heading.group(2)))
            index += 1
            continue
        first = _cells(line)
        separator = _cells(lines[index + 1]) if index + 1 < len(lines) else None
        if (first and separator and len(first) == len(separator)
                and all(_SEPARATOR.fullmatch(cell) for cell in separator)):
            if len(first) > _MAX_TABLE_COLUMNS:
                raise ValueError('The document table exceeds the column limit.')
            rows = [first]
            index += 2
            while index < len(lines):
                row = _cells(lines[index])
                if row is None or len(row) != len(first):
                    break
                rows.append(row)
                if len(rows) > _MAX_TABLE_ROWS:
                    raise ValueError('The document table exceeds the row limit.')
                index += 1
            _table(document, rows)
            continue
        item = _LIST.match(line)
        if item:
            style = 'List Bullet' if item.group('bullet') else 'List Number'
            paragraph = document.add_paragraph(style=style)
            if item.group('number'):
                _number(document, paragraph, int(item.group('number')))
            _inline(paragraph, item.group('text'))
        elif line.startswith('> '):
            _inline(document.add_paragraph(style='Quote'), line[2:])
        else:
            _inline(document.add_paragraph(), line)
        index += 1
    return document, _plain_text(document)

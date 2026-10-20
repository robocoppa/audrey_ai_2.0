"""Local-only document preparation for the internal protected upload path.

Prepared vectors belong to an isolated collection; ordinary user collections
never contain protected material, even before the owned upload commits.
"""

from __future__ import annotations

import asyncio
import uuid
import zipfile
from contextlib import asynccontextmanager
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import magic

from audrey.kb.chunk import chunk_text
from audrey.kb.embed import text_embedder
from audrey.kb.extract import SUFFIX_MIMES
from audrey.kb.qdrant import build_text_point, normalize_source
from audrey.kb.user_store import sanitize_user
from audrey.pipeline.summarise import (
    DEFAULT_INPUT_BUDGET,
    DEFAULT_MODEL,
    DOCUMENT_SUMMARY_SYSTEM,
    SUMMARY_MAX_OUTPUT_TOKENS,
    _document_units,
    brief_video_summary,
    build_input,
)
from audrey.processing_policy import ProcessingPolicyError, current_processing_policy

_MAX_TEXT_CHARS = 2_000_000
_MAX_PAGES = 2_000
_MAX_EXPANDED_BYTES = 64 * 1024 * 1024
_PLAIN_SUFFIXES = frozenset({".txt", ".log", ".md", ".rst", ".csv"})
_DOCUMENT_SUFFIXES = _PLAIN_SUFFIXES | {".pdf", ".docx", ".xlsx", ".html", ".htm"}


class PrivateDocumentError(RuntimeError):
    """Content-free refusal for malformed, unsupported or oversized documents."""

    def __init__(self, code: str = "private_document_invalid") -> None:
        self.code = code if code in {
            "private_document_invalid", "private_document_unsupported",
            "private_document_requires_ocr", "private_document_limit",
        } else "private_document_invalid"
        messages = {
            "private_document_invalid": "The document could not be read completely.",
            "private_document_unsupported": "This file type is not supported for Private processing.",
            "private_document_requires_ocr": "This PDF needs OCR, which is not available for Private processing yet.",
            "private_document_limit": "The document exceeds the supported Private processing limits.",
        }
        super().__init__(messages[self.code])


def private_text_collection(user: str) -> str:
    return f"kb_private_text_{sanitize_user(user)}"


def _bounded(text: str) -> str:
    if len(text) > _MAX_TEXT_CHARS:
        raise PrivateDocumentError("private_document_limit")
    if not text.strip():
        raise PrivateDocumentError()
    return text


def _check_package(path: Path) -> None:
    with zipfile.ZipFile(path) as archive:
        if sum(entry.file_size for entry in archive.infolist()) > _MAX_EXPANDED_BYTES:
            raise PrivateDocumentError("private_document_limit")


def extract_private_document(path: Path) -> tuple[str, str]:
    """Read supported formats locally, refusing partial PDF/page extraction.

    Sniff failures are refused, without the ordinary suffix fallback or parser
    exception logging. Painted pages without a text layer need the later OCR
    path. Provably blank pages retain their numbered marker without inventing
    extracted content or silently dropping a page.
    """
    try:
        suffix = path.suffix.lower()
        if suffix not in _DOCUMENT_SUFFIXES:
            raise PrivateDocumentError("private_document_unsupported")
        sniffed = magic.from_file(str(path), mime=True)
        expected = SUFFIX_MIMES[suffix]
        compatible = {expected}
        if suffix in _PLAIN_SUFFIXES:
            compatible |= {"text/plain", "text/csv", "text/markdown", "text/x-markdown", "text/x-rst"}
        if suffix in {".docx", ".xlsx"}:
            compatible |= {"application/zip"}
        if sniffed not in compatible:
            raise PrivateDocumentError("private_document_unsupported")
        if suffix in _PLAIN_SUFFIXES:
            return expected, _bounded(path.read_text(encoding="utf-8", errors="strict"))
        if suffix == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(str(path), strict=True)
            if reader.is_encrypted and not reader.decrypt(""):
                raise PrivateDocumentError()
            if not 0 < len(reader.pages) <= _MAX_PAGES:
                raise PrivateDocumentError("private_document_limit")
            parts: list[str] = []
            count = 0
            for index, page in enumerate(reader.pages, 1):
                text = page.extract_text()
                if not isinstance(text, str) or not text.strip():
                    contents = page.get_contents()
                    if contents is not None and contents.get_data().strip():
                        raise PrivateDocumentError("private_document_requires_ocr")
                    text = ""
                count += len(text)
                if count > _MAX_TEXT_CHARS:
                    raise PrivateDocumentError("private_document_limit")
                parts.append(f"--- Page {index} ---\n{text}")
            if count == 0:
                raise PrivateDocumentError()
            return expected, _bounded("\n\n".join(parts))
        if suffix == ".docx":
            from docx import Document
            _check_package(path)
            document = Document(str(path))
            parts = [paragraph.text for paragraph in document.paragraphs]
            parts.extend("\t".join(cell.text for cell in row.cells)
                         for table in document.tables for row in table.rows)
            return expected, _bounded("\n".join(parts))
        if suffix == ".xlsx":
            from openpyxl import load_workbook
            _check_package(path)
            book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
            try:
                if len(book.worksheets) > 20:
                    raise PrivateDocumentError("private_document_limit")
                parts = []
                count = 0
                for sheet in book.worksheets:
                    if sheet.max_row > 50_000 or sheet.max_column > 100:
                        raise PrivateDocumentError("private_document_limit")
                    parts.append(f"Sheet: {sheet.title} (formulas are not evaluated)")
                    for row in sheet.iter_rows(values_only=True):
                        line = "\t".join("" if cell is None else str(cell) for cell in row)
                        count += len(line) + 1
                        if count > _MAX_TEXT_CHARS:
                            raise PrivateDocumentError("private_document_limit")
                        parts.append(line)
                return expected, _bounded("\n".join(parts))
            finally:
                book.close()
        from bs4 import BeautifulSoup
        raw = _bounded(path.read_text(encoding="utf-8", errors="strict"))
        soup = BeautifulSoup(raw, "html.parser")
        for tag in soup(["script", "style", "noscript"]):
            tag.decompose()
        return expected, _bounded(soup.get_text("\n", strip=True))
    except PrivateDocumentError:
        raise
    except Exception:  # noqa: BLE001 - redact untrusted parser diagnostics
        raise PrivateDocumentError() from None


@dataclass(frozen=True, slots=True)
class PreparedPrivateDocument:
    mime: str
    text: str
    summary: str
    points: list[Any]
    dimensions: int


async def prepare_private_document(
    app: Any, path: Path, *, user: str, file_id: str, filename: str,
    source_bytes: int, stamp: str,
) -> PreparedPrivateDocument:
    """Prepare all summaries/vectors before persistence, within a verified scope."""
    policy = current_processing_policy()
    if policy is None:
        raise ProcessingPolicyError("private_processing_unavailable")
    embedder = text_embedder(app.state.ollama, app.state.cfg.raw)
    policy.assert_model_allowed(embedder.model)
    if policy.local_models[embedder.model] != embedder.profile.digest:
        raise ProcessingPolicyError("local_model_digest_mismatch")
    mime, text = await asyncio.to_thread(extract_private_document, path)
    summary = ""
    settings = ((app.state.cfg.raw.get("kb") or {}).get("video") or {})
    summary_model = str(settings.get("summarise_model") or DEFAULT_MODEL).strip()
    # A summary is optional; never widen the verified roster or fall back.
    if mime == "application/pdf":
        try:
            policy.assert_model_allowed(summary_model)
        except ProcessingPolicyError:
            pass
        else:
            material = build_input(
                [{"text": unit} for unit in _document_units(text)], [],
                budget=int(settings.get("summary_input_chars", DEFAULT_INPUT_BUDGET)),
                media_kind="document",
            )
            try:
                caps = await app.state.ollama.capabilities(summary_model)
                async with app.state.gate.acquire(summary_model, location="local", user_id=user):
                    response = await app.state.ollama.chat(
                        model=summary_model,
                        messages=[{"role": "system", "content": DOCUMENT_SUMMARY_SYSTEM},
                                  {"role": "user", "content": material}],
                        think=False if "thinking" in caps else None,
                        options={"num_predict": SUMMARY_MAX_OUTPUT_TOKENS},
                        timeout_s=float(settings.get("summary_timeout_s", 180)),
                    )
                summary = brief_video_summary(str((response.get("message") or {}).get("content") or ""))
            except ProcessingPolicyError:
                raise
            except Exception:  # noqa: BLE001 - optional summary; redact provider diagnostics
                # The complete document remains useful without a description.
                # Do not log provider errors, names, text or paths.
                summary = ""
    chunks = await asyncio.to_thread(chunk_text, text)
    if not chunks:
        raise PrivateDocumentError()
    vectors = await embedder.embed_many([chunk.text for chunk in chunks] + ([summary] if summary else []))
    extras = {
        "user": user, "file_id": file_id, "filename": filename,
        "mime": mime, "bytes": source_bytes, "uploaded_at": stamp,
        "private_required": True, **embedder.profile.provenance,
    }
    stat = await asyncio.to_thread(path.stat)
    points = [build_text_point(
        source=normalize_source(path), chunk_idx=chunk.idx, text=chunk.text,
        vector=vector, mtime=stat.st_mtime, sparse=True,
        extra={**extras, "artifact": "text"},
    ) for chunk, vector in zip(chunks, vectors[:len(chunks)], strict=True)]
    if summary:
        points.append(build_text_point(
            source=normalize_source(path.with_suffix(".summary.txt")), chunk_idx=0,
            text=summary, vector=vectors[-1], mtime=stat.st_mtime, sparse=True,
            extra={**extras, "artifact": "summary"},
        ))
    return PreparedPrivateDocument(mime, text, summary, points, embedder.profile.dimensions)


@dataclass(frozen=True, slots=True)
class PrivateDocumentSource:
    file_id: str
    filename: str
    mime: str
    source_bytes: int
    path: Path


@dataclass(frozen=True, slots=True)
class PrivateDocumentText:
    file_id: str
    filename: str
    mime: str
    source_bytes: int
    text: str


@asynccontextmanager
async def private_document_source(app: Any, principal: Any, file_id: str):
    """Hold the file operation lock across an owner-authorized local read."""
    if current_processing_policy() is None:
        raise ProcessingPolicyError("private_processing_unavailable")
    try:
        if str(uuid.UUID(file_id)) != file_id:
            raise ValueError
        user = principal.storage_namespace
        locks = app.state.file_operation_locks
        db = app.state.uploads_db
    except (AttributeError, ValueError):
        raise ProcessingPolicyError("private_processing_unavailable") from None
    async with locks.hold(file_id):
        row = await db.get_upload(file_id)
        if (
            not row or row.get("user") != user or not row.get("private_required")
            or row.get("status") != "ready" or row.get("kind") != "text"
            or row.get("collection") != private_text_collection(user)
            or row.get("source_freed_at")
        ):
            raise ProcessingPolicyError("private_processing_unavailable")
        root = Path((app.state.cfg.raw.get("kb") or {}).get("upload_root", "/data/uploads"))
        path = root / sanitize_user(user) / f"{file_id}{Path(row['filename']).suffix.lower()}"
        try:
            if path.is_symlink() or path.resolve() != path.absolute() or path.stat().st_size != row["bytes"]:
                raise ValueError
        except (OSError, ValueError):
            raise ProcessingPolicyError("private_processing_unavailable") from None
        yield PrivateDocumentSource(
            file_id=file_id, filename=str(row["filename"]), mime=str(row["mime"]),
            source_bytes=int(row["bytes"]), path=path,
        )


async def read_private_document(app: Any, principal: Any, file_id: str) -> PrivateDocumentText:
    """Read the exact locally extracted text; never reparse, embed or call a model."""
    async with private_document_source(app, principal, file_id) as source:
        path = source.path.with_suffix(".ocr.txt")
        try:
            if path.is_symlink() or path.resolve() != path.absolute() or path.stat().st_size > _MAX_TEXT_CHARS * 4:
                raise ValueError
            text = _bounded(await asyncio.to_thread(path.read_text, encoding="utf-8"))
        except Exception:  # noqa: BLE001 - never disclose paths or document text
            raise ProcessingPolicyError("private_processing_unavailable") from None
        return PrivateDocumentText(
            source.file_id, source.filename, source.mime, source.source_bytes, text,
        )

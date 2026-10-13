"""Bounded read-only spreadsheet analysis; the owning route resolves paths.

First row is the header. Cells are text/null; decimal totals are returned as
strings. No SQL, formulas or user code run. Oversized tables fail before totals.
"""
from __future__ import annotations

import csv
import io
import json
import re
import zipfile
from collections import Counter
from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation, localcontext
from pathlib import Path
from typing import Any, Literal

from pydantic import BaseModel, ConfigDict, Field, model_validator

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
TABLE_MIMES = frozenset({"text/csv", XLSX_MIME})
MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 64 * 1024 * 1024
MAX_ROWS = 50_000
MAX_COLUMNS = 100
MAX_CELLS = 1_000_000
MAX_CELL_CHARS = 8_192
MAX_GROUPS = 1_000
MAX_SHEETS = 20
MAX_RESULT_CHARS = 2_500
MAX_RESULT_BYTES = 2_500
MAX_HEADER_CHARS = 128
DISPLAY_CELL_CHARS = 500
_NUMBER = re.compile(r"[+-]?(?:[0-9]+(?:\.[0-9]*)?|\.[0-9]+)(?:[eE]([+-]?[0-9]+))?\Z")
_COORDINATE = re.compile(r"([A-Z]+)([1-9][0-9]*)\Z")
_XML_NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"


class TableAnalysisError(ValueError):
    """Safe explanation of invalid content, unavailable readers, or a limit."""


def _number(value: str | None) -> Decimal | None:
    if value is None:
        return None
    text = value.strip()
    match = _NUMBER.fullmatch(text)
    if len(text) > 128 or match is None:
        return None
    if match.group(1) is not None and abs(int(match.group(1))) > 100:
        return None
    try:
        number = Decimal(text)
    except InvalidOperation:
        return None
    return number if number.is_finite() else None


class TableFilter(BaseModel):
    model_config = ConfigDict(extra="forbid")
    column: str = Field(min_length=1, max_length=MAX_CELL_CHARS)
    op: Literal["eq", "ne", "contains", "gt", "gte", "lt", "lte", "is_blank", "not_blank"]
    value: str | None = Field(default=None, max_length=MAX_CELL_CHARS)

    @model_validator(mode="after")
    def validate_value(self) -> TableFilter:
        if self.op not in {"is_blank", "not_blank"} and self.value is None:
            raise ValueError("a filter value is required")
        if self.op in {"gt", "gte", "lt", "lte"} and _number(self.value) is None:
            raise ValueError("numeric filters require a plain finite number")
        return self


class TableAggregation(BaseModel):
    model_config = ConfigDict(extra="forbid")
    op: Literal["count", "sum", "mean", "min", "max"]
    column: str | None = Field(default=None, min_length=1, max_length=MAX_CELL_CHARS)

    @model_validator(mode="after")
    def validate_column(self) -> TableAggregation:
        if self.op != "count" and self.column is None:
            raise ValueError("numeric aggregations require a column")
        return self


class TableAnalysisRequest(BaseModel):
    model_config = ConfigDict(extra="forbid")
    action: Literal["inspect", "rows", "aggregate", "duplicates"] = "inspect"
    sheet: str | None = Field(default=None, min_length=1, max_length=256)
    columns: list[str] = Field(default_factory=list, max_length=MAX_COLUMNS)
    filters: list[TableFilter] = Field(default_factory=list, max_length=8)
    group_by: list[str] = Field(default_factory=list, max_length=4)
    aggregations: list[TableAggregation] = Field(default_factory=list, max_length=10)
    column_offset: int = Field(default=0, ge=0, le=MAX_COLUMNS)
    column_limit: int = Field(default=10, ge=1, le=MAX_COLUMNS)
    offset: int = Field(default=0, ge=0, le=MAX_ROWS)
    limit: int = Field(default=10, ge=1, le=100)

    @model_validator(mode="after")
    def validate_action(self) -> TableAnalysisRequest:
        if self.action == "aggregate" and not self.aggregations:
            raise ValueError("aggregate requires at least one aggregation")
        if self.action != "aggregate" and (self.aggregations or self.group_by):
            raise ValueError("aggregations and group_by apply only to aggregate")
        if self.action == "inspect" and self.filters:
            raise ValueError("use rows for filtered results")
        for values in (self.columns, self.group_by):
            if len(values) != len(set(values)):
                raise ValueError("column selections must be unique")
        keys = [(agg.op, agg.column) for agg in self.aggregations]
        if len(keys) != len(set(keys)):
            raise ValueError("aggregations must be unique")
        return self


@dataclass
class _Table:
    sheets: list[str]
    sheet: str
    headers: list[str]
    original_headers: list[str | None]
    rows: list[list[str | None]]
    formula_cells: int = 0
    formulas_without_cache: int = 0


def _cell(value: Any) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        text = "true" if value else "false"
    elif isinstance(value, (datetime, date, time)):
        text = value.isoformat()
    else:
        text = str(value)
    if len(text) > MAX_CELL_CHARS:
        raise TableAnalysisError(f"a cell exceeds the {MAX_CELL_CHARS}-character limit")
    return text


def _headers(raw: list[str | None]) -> list[str]:
    result: list[str] = []
    used: set[str] = set()
    for index, value in enumerate(raw):
        base = (value or "").strip() or f"column_{index + 1}"
        if len(base) > MAX_HEADER_CHARS:
            raise TableAnalysisError(f"a column header exceeds the {MAX_HEADER_CHARS}-character limit")
        name, suffix = base, 2
        while name in used:
            name = f"{base} ({suffix})"
            suffix += 1
        used.add(name)
        result.append(name)
    return result


def _finish(
    sheets: list[str], sheet: str, raw_rows: list[list[str | None]],
    formulas: int = 0, uncached: int = 0,
) -> _Table:
    if not raw_rows:
        raise TableAnalysisError("the table is empty; the first row must contain headers")
    width = max(len(row) for row in raw_rows)
    if not width:
        raise TableAnalysisError("the table has no columns")
    if width > MAX_COLUMNS or (len(raw_rows) - 1) * width > MAX_CELLS:
        raise TableAnalysisError("the table exceeds the column or cell limit")
    original = raw_rows[0] + [None] * (width - len(raw_rows[0]))
    rows = [row + [None] * (width - len(row)) for row in raw_rows[1:]]
    return _Table(sheets, sheet, _headers(original), original, rows, formulas, uncached)


def _read_csv(data: bytes, sheet: str | None) -> _Table:
    if sheet not in {None, "CSV"}:
        raise TableAnalysisError("CSV has one table named CSV")
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise TableAnalysisError("CSV must use UTF-8 encoding") from error
    if "\x00" in text:
        raise TableAnalysisError("CSV contains invalid null bytes")
    try:
        dialect = csv.Sniffer().sniff(text[:8_192], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows: list[list[str | None]] = []
    try:
        for raw in csv.reader(io.StringIO(text, newline=""), dialect, strict=True):
            if not raw:
                continue
            if len(raw) > MAX_COLUMNS:
                raise TableAnalysisError(f"the table exceeds the {MAX_COLUMNS}-column limit")
            if len(rows) > MAX_ROWS:
                raise TableAnalysisError(f"the table exceeds the {MAX_ROWS}-row limit")
            rows.append([_cell(value) for value in raw])
    except csv.Error as error:
        raise TableAnalysisError("CSV cannot be parsed") from error
    return _finish(["CSV"], "CSV", rows)


def _xlsx_preflight(data: bytes) -> dict[str, tuple[int, int, int, int]]:
    try:
        from defusedxml.ElementTree import iterparse
    except ImportError as error:
        raise TableAnalysisError("Excel analysis requires the protected XML reader") from error
    dimensions: dict[str, tuple[int, int, int, int]] = {}
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > 1_000 or sum(item.file_size for item in entries) > MAX_UNCOMPRESSED_BYTES:
                raise TableAnalysisError("Excel contents exceed the archive limit")
            names = [item.filename for item in entries]
            if len(names) != len(set(names)):
                raise TableAnalysisError("Excel archive contains duplicate entries")
            worksheets = [
                name for name in names
                if name.startswith("xl/worksheets/") and name.endswith(".xml")
                and "/" not in name.removeprefix("xl/worksheets/")
            ]
            if len(worksheets) > MAX_SHEETS:
                raise TableAnalysisError(f"the workbook exceeds the {MAX_SHEETS}-sheet limit")
            for name in worksheets:
                max_row = max_col = formulas = uncached = cells = 0
                with archive.open(name) as source:
                    for _, element in iterparse(source, events=("end",)):
                        if element.tag == _XML_NS + "c":
                            coord = _COORDINATE.fullmatch(element.attrib.get("r", ""))
                            if coord is None:
                                raise TableAnalysisError("Excel contains an invalid cell coordinate")
                            row_number, column_number = int(coord.group(2)), 0
                            for letter in coord.group(1):
                                column_number = column_number * 26 + ord(letter) - ord("A") + 1
                            max_row, max_col = max(max_row, row_number), max(max_col, column_number)
                            cells += 1
                            if max_row > MAX_ROWS + 1 or max_col > MAX_COLUMNS or cells > MAX_CELLS:
                                raise TableAnalysisError("Excel sheet exceeds row, column or cell limits")
                            if element.find(_XML_NS + "f") is not None:
                                formulas += 1
                                cached = element.find(_XML_NS + "v")
                                if cached is None or cached.text in {None, ""}:
                                    uncached += 1
                            element.clear()
                        elif element.tag == _XML_NS + "row":
                            element.clear()
                dimensions["/" + name] = (max_row, max_col, formulas, uncached)
    except TableAnalysisError:
        raise
    except Exception as error:  # Safe failure for untrusted archive/XML parsers.
        raise TableAnalysisError("Excel archive cannot be parsed safely") from error
    return dimensions


def _read_xlsx(data: bytes, selected: str | None) -> _Table:
    dimensions = _xlsx_preflight(data)
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise TableAnalysisError("Excel analysis is unavailable; install the Excel reader") from error
    workbook = None
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
        if any(len(name) > 31 for name in workbook.sheetnames):
            raise TableAnalysisError("Excel sheet names exceed the 31-character analysis limit")
        if not workbook.sheetnames:
            raise TableAnalysisError("the workbook has no sheets")
        if len(workbook.sheetnames) > MAX_SHEETS:
            raise TableAnalysisError(f"the workbook exceeds the {MAX_SHEETS}-sheet limit")
        sheet_name = selected or workbook.sheetnames[0]
        if selected is None:
            # A common workbook starts with an unused blank tab. Default to
            # the first tab with cells; still return all sheet names and the
            # chosen tab explicitly so later requests can select another.
            for candidate in workbook.sheetnames:
                candidate_path = "/" + str(getattr(workbook[candidate], "_worksheet_path", "")).lstrip("/")
                bounds = dimensions.get(candidate_path)
                if bounds and bounds[0] and bounds[1]:
                    sheet_name = candidate
                    break
        if sheet_name not in workbook.sheetnames:
            raise TableAnalysisError("the requested sheet does not exist")
        worksheet = workbook[sheet_name]
        worksheet_path = getattr(worksheet, "_worksheet_path", "")
        if not worksheet_path.startswith("/"):
            worksheet_path = "/" + worksheet_path
        if worksheet_path not in dimensions:
            raise TableAnalysisError("Excel sheet could not be bounded")
        row_count, column_count, formulas, uncached = dimensions[worksheet_path]
        if not row_count or not column_count:
            raise TableAnalysisError("the sheet is empty")
        if row_count * column_count > MAX_CELLS:
            raise TableAnalysisError(f"the table exceeds the {MAX_CELLS}-cell limit")
        # Untrusted dimension metadata can under-report the actual used area.
        worksheet.reset_dimensions()
        rows = [
            [_cell(value) for value in row]
            for row in worksheet.iter_rows(
                min_row=1, max_row=row_count, max_col=column_count, values_only=True,
            )
        ]
        return _finish(workbook.sheetnames, sheet_name, rows, formulas, uncached)
    except TableAnalysisError:
        raise
    except Exception as error:  # Safe failure for untrusted workbook parsing.
        raise TableAnalysisError("Excel workbook cannot be read") from error
    finally:
        if workbook is not None:
            workbook.close()


def _load(path: Path, mime: str, sheet: str | None) -> _Table:
    if mime not in TABLE_MIMES:
        raise TableAnalysisError("analysis supports only CSV and .xlsx files")
    try:
        if path.is_symlink() or not path.is_file():
            raise TableAnalysisError("the original file is unavailable")
        if path.stat().st_size > MAX_FILE_BYTES:
            raise TableAnalysisError(f"the file exceeds the {MAX_FILE_BYTES}-byte analysis limit")
        with path.open("rb") as source:
            data = source.read(MAX_FILE_BYTES + 1)
    except OSError as error:
        raise TableAnalysisError("the original file is unavailable") from error
    if len(data) > MAX_FILE_BYTES:
        raise TableAnalysisError("the file exceeds the analysis limit")
    return _read_csv(data, sheet) if mime == "text/csv" else _read_xlsx(data, sheet)


def _positions(table: _Table, names: list[str]) -> list[int]:
    unknown = [name for name in names if name not in table.headers]
    if unknown:
        raise TableAnalysisError(f"unknown column: {unknown[0]}; inspect the table for exact names")
    return [table.headers.index(name) for name in names]


def _matches(value: str | None, condition: TableFilter) -> bool:
    blank = value is None or not value.strip()
    if condition.op == "is_blank":
        return blank
    if condition.op == "not_blank":
        return not blank
    if condition.op == "eq":
        return value == condition.value
    if condition.op == "ne":
        return value != condition.value
    if condition.op == "contains":
        return value is not None and condition.value is not None and condition.value in value
    left, right = _number(value), _number(condition.value)
    if left is None or right is None:
        return False
    return {"gt": left > right, "gte": left >= right, "lt": left < right, "lte": left <= right}[condition.op]


def _display(rows: list[list[str | None]]) -> tuple[list[list[str | None]], bool]:
    truncated = any(
        value is not None and len(value) > DISPLAY_CELL_CHARS for row in rows for value in row
    )
    return [
        [value[:DISPLAY_CELL_CHARS] if value is not None else None for value in row]
        for row in rows
    ], truncated


def _measure(rows: list[list[str | None]], column: int | None, op: str) -> dict[str, Any]:
    values = [row[column] for row in rows] if column is not None else []
    missing = sum(value is None or not value.strip() for value in values)
    numeric = [number for value in values if (number := _number(value)) is not None]
    result: dict[str, Any] = {"op": op, "rows": len(rows)}
    if column is not None:
        result.update({
            "missing_cells": missing, "numeric_cells": len(numeric),
            "non_numeric_cells": len(values) - len(numeric) - missing,
        })
    if op == "count":
        result["value"] = len(rows) if column is None else len(values) - missing
        result["counts"] = "rows" if column is None else "nonblank cells"
    elif not numeric:
        result["value"] = None
    else:
        with localcontext() as context:
            context.prec = 512
            total = sum(numeric, Decimal(0))
            if op == "sum":
                value = total
            elif op == "mean":
                context.prec = 28
                value = total / len(numeric)
                result["precision"] = "28 significant decimal digits"
            elif op == "min":
                value = min(numeric)
            else:
                value = max(numeric)
        result["value"] = format(value, "f")
    return result


def analyze_table(
    path: Path, mime: str, request: TableAnalysisRequest | dict[str, Any],
) -> dict[str, Any]:
    """Analyze an owner-resolved original; the caller supplies identity/path gates."""
    spec = request if isinstance(request, TableAnalysisRequest) else TableAnalysisRequest.model_validate(request)
    table = _load(path, mime, spec.sheet)
    if spec.action == "inspect" and not spec.columns:
        columns = table.headers[spec.column_offset:spec.column_offset + spec.column_limit]
    else:
        columns = spec.columns or table.headers
    positions = _positions(table, columns)
    filter_positions = _positions(table, [item.column for item in spec.filters])
    filtered = [
        row for row in table.rows
        if all(_matches(row[index], item) for index, item in zip(filter_positions, spec.filters, strict=True))
    ]
    result: dict[str, Any] = {
        "action": spec.action, "sheet": table.sheet, "sheets": table.sheets,
        "column_count": len(table.headers),
        "total_rows": len(table.rows), "matched_rows": len(filtered),
        "full_table_scanned": True,
        "semantics": "First row is headers. Numeric operations use plain finite numbers; currency/percent text is excluded. Blank/non-numeric exclusions are counted. Excel formulas use saved cached values only; never executed.",
        "formula_cells": table.formula_cells,
        "formulas_without_cached_values": table.formulas_without_cache,
    }
    if spec.action == "aggregate":
        group_positions = _positions(table, spec.group_by)
        aggregation_positions = [
            _positions(table, [agg.column])[0] if agg.column is not None else None
            for agg in spec.aggregations
        ]
        groups: dict[tuple[str | None, ...], list[list[str | None]]] = {}
        for row in filtered:
            key = tuple(row[index] for index in group_positions)
            if key not in groups and len(groups) >= MAX_GROUPS:
                raise TableAnalysisError(f"the result exceeds the {MAX_GROUPS}-group limit")
            groups.setdefault(key, []).append(row)
        if not group_positions and not groups:
            groups[()] = []
        result["group_by"], result["groups_total"] = spec.group_by, len(groups)
        items = list(groups.items())[spec.offset:spec.offset + spec.limit]
        result["groups"] = [
            {
                "key": list(key), "rows": len(rows),
                "aggregations": [
                    {"column": agg.column, **_measure(rows, column, agg.op)}
                    for agg, column in zip(spec.aggregations, aggregation_positions, strict=True)
                ],
            }
            for key, rows in items
        ]
        result["has_more"] = spec.offset + len(items) < len(groups)
    elif spec.action == "duplicates":
        counts = Counter(tuple(row[index] for index in positions) for row in filtered)
        duplicates = [(key, count) for key, count in counts.items() if count > 1]
        page = duplicates[spec.offset:spec.offset + spec.limit]
        displayed, truncated = _display([list(key) for key, _ in page])
        result.update({
            "selected_columns": columns, "duplicate_groups": len(duplicates),
            "extra_duplicate_rows": sum(count - 1 for _, count in duplicates),
            "duplicates": [
                {"values": values, "occurrences": count}
                for values, (_, count) in zip(displayed, page, strict=True)
            ],
            "has_more": spec.offset + len(page) < len(duplicates),
            "display_cells_truncated": truncated,
        })
    else:
        page = filtered[spec.offset:spec.offset + spec.limit]
        displayed, truncated = _display([[row[index] for index in positions] for row in page])
        result.update({
            "selected_columns": columns, "rows": displayed, "offset": spec.offset,
            "has_more": spec.offset + len(page) < len(filtered),
            "display_cells_truncated": truncated,
        })
        if spec.action == "inspect":
            result["columns"] = columns
            result["column_offset"] = spec.column_offset
            result["next_column_offset"] = (
                spec.column_offset + len(columns)
                if not spec.columns and spec.column_offset + len(columns) < len(table.headers)
                else None
            )
            result["original_headers"] = [table.original_headers[index] for index in positions]
            result["column_stats"] = [
                {
                    "column": name,
                    "blank_cells": sum(row[index] is None or not row[index].strip() for row in table.rows),
                    "numeric_cells": sum(_number(row[index]) is not None for row in table.rows),
                }
                for index, name in zip(positions, columns, strict=True)
            ]
    # For inspect, shrink the column page before paging samples. Exact names
    # remain available through next_column_offset; none are silently ellipsized.
    def metadata_fits() -> bool:
        metadata = {key: value for key, value in result.items() if key != "rows"}
        metadata["rows"] = result.get("rows", [])[:1]
        encoded = json.dumps(metadata, ensure_ascii=False)
        return len(encoded) + 100 <= MAX_RESULT_CHARS and len(encoded.encode("utf-8")) + 100 <= MAX_RESULT_BYTES

    if spec.action == "inspect" and not spec.columns:
        while len(result["columns"]) > 1 and not metadata_fits():
            count = len(result["columns"]) - 1
            result["columns"] = result["columns"][:count]
            result["selected_columns"] = result["selected_columns"][:count]
            result["original_headers"] = result["original_headers"][:count]
            result["column_stats"] = result["column_stats"][:count]
            result["rows"] = [row[:count] for row in result["rows"]]
            result["next_column_offset"] = spec.column_offset + count

    # Fit a complete JSON page to both character and transport-byte budgets.
    # Totals/aggregates are computed from the full bounded table before paging.
    field = "groups" if spec.action == "aggregate" else "duplicates" if spec.action == "duplicates" else "rows"
    total_items = result["groups_total"] if spec.action == "aggregate" else result["duplicate_groups"] if spec.action == "duplicates" else len(filtered)
    page_items = result[field]
    requested_count = len(page_items)

    def set_page(count: int) -> bool:
        result[field] = page_items[:count]
        result["returned_items"] = count
        result["offset"] = spec.offset
        result["has_more"] = spec.offset + count < total_items
        result["next_offset"] = spec.offset + count if result["has_more"] else None
        encoded = json.dumps(result, ensure_ascii=False)
        return len(encoded) <= MAX_RESULT_CHARS and len(encoded.encode("utf-8")) <= MAX_RESULT_BYTES

    if not set_page(requested_count):
        low, high = 0, requested_count
        while low < high:
            middle = (low + high + 1) // 2
            if set_page(middle):
                low = middle
            else:
                high = middle - 1
        if not set_page(low) or (low == 0 and requested_count):
            raise TableAnalysisError("the result is too large; request fewer columns or a smaller limit")
    return result


def extract_table_text(path: Path, mime: str, max_chars: int = 500_000) -> str:
    """Build an explicitly bounded first-sheet preview for upload indexing."""
    if max_chars < 128 or max_chars > 1_000_000:
        raise ValueError("preview length must be between 128 and 1000000")
    table = _load(path, mime, None)
    lines = [
        f"Sheet: {table.sheet}; sheets: {', '.join(table.sheets)}",
        f"Rows: {len(table.rows)}. First row is the header. Formulas are not evaluated.",
        "\t".join(table.headers),
    ]
    used, truncated = sum(len(line) + 1 for line in lines), False
    for row in table.rows:
        line = "\t".join(value or "" for value in row)
        if used + len(line) + 1 > max_chars:
            truncated = True
            break
        lines.append(line)
        used += len(line) + 1
    if truncated or len(table.sheets) > 1:
        lines.append(
            "[Preview only. Other sheets or remaining rows may be omitted. "
            "Use the spreadsheet analysis tool for complete bounded analysis.]"
        )
    return "\n".join(lines)


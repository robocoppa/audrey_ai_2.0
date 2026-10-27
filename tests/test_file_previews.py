"""Browser projections preserve actual generated contents without downloading."""
from __future__ import annotations

import io
import zipfile

import pytest
from openpyxl import Workbook
from test_generated_files import _create, _run
from test_generated_files import output_state as output_state

from audrey.file_previews import MAX_ROWS, FilePreviewError, workbook_preview
from audrey.processing_policy import ProcessingPolicy, processing_scope

CASES = [
    ('txt', {'content': 'Meeting notes\nNext step: follow up.'}, 'text'),
    ('md', {'content': '# Weekend plan\n\n- Read\n- Walk'}, 'text'),
    ('json', {'data': {'ready': True, 'total': 65}}, 'text'),
    ('docx', {'content': 'Weekend plan\n\nRead a book.\n\nTake a walk.'}, 'document'),
    ('csv', {'rows': [['Item', 'Cost'], ['Food', 40], ['Travel', 25]]}, 'table'),
    ('xlsx', {'sheets': [
        {'name': 'Budget', 'rows': [['Item', 'Cost'], ['Food', 40], ['Travel', 25]]},
        {'name': 'Notes', 'rows': [['=NOT_A_FORMULA', '<script>alert(1)</script>'], [None, True]]},
    ]}, 'table'),
]


@pytest.mark.parametrize('private', [False, True])
@pytest.mark.parametrize('format, fields, kind', CASES)
async def test_preview_all_generated_formats_without_changing_file_or_case(output_state, private, format, fields, kind):
    state = output_state
    conversation, run = await _run(state, private=private)
    policy = ProcessingPolicy(local_models={'local:latest': 'a' * 64}, ollama_base_url='http://ollama:11434',
        allowed_tools=frozenset({'create_file'}), tool_servers=frozenset({'http://ollama:11434'})) if private else None
    with processing_scope(policy):
        item = await _create(state, conversation, run, filename=f'weekend.{format}', format=format, **fields)
    assert 'error' not in item, item
    before = await state.db.get_upload(item['file_id'])
    path = f"/api/files/{item['file_id']}/preview"
    response = await state.client.get(path)
    assert response.status_code == 200, response.text
    preview = response.json()
    assert preview['id'] == item['file_id'] and preview['filename'] == f'weekend.{format}'
    assert preview['format'] == format and preview['kind'] == kind
    assert preview['truncated'] is False
    if format in ('txt', 'md'):
        assert preview['text'] == fields['content']
    elif format == 'docx':
        assert preview['paragraphs'] == ['Weekend plan', 'Read a book.', 'Take a walk.']
    elif format == 'json':
        import json
        assert json.loads(preview['text']) == fields['data']
    elif format == 'csv':
        assert preview['sheets'][0]['rows'] == [['Item', 'Cost'], ['Food', '40'], ['Travel', '25']]
    else:
        assert [sheet['name'] for sheet in preview['sheets']] == ['Budget', 'Notes']
        assert preview['sheets'][0]['rows'] == fields['sheets'][0]['rows']
        assert preview['sheets'][1]['rows'] == fields['sheets'][1]['rows']
    assert response.headers['cache-control'] == 'private, no-store'
    assert response.headers['x-content-type-options'] == 'nosniff'
    assert 'content-disposition' not in response.headers
    assert await state.db.get_upload(item['file_id']) == before
    assert (await state.client.get(item['download_url'])).status_code == 200
    state.current[0] = state.other
    assert (await state.client.get(path)).status_code == 404
    state.current[0] = state.owner
    assert (await state.client.delete(f"/api/files/{item['file_id']}")).status_code == 200
    assert (await state.client.get(path)).status_code == 404


async def test_larger_preview_honestly_reports_cutoff(output_state):
    state = output_state
    conversation, run = await _run(state)
    item = await _create(state, conversation, run, filename='long.csv', format='csv', rows=[[i] for i in range(MAX_ROWS + 5)])
    response = await state.client.get(f"/api/files/{item['file_id']}/preview")
    assert response.status_code == 200, response.text
    result = response.json()
    assert result['truncated'] and len(result['sheets'][0]['rows']) == MAX_ROWS
    assert result['sheets'][0]['truncated']
    assert 'part of this file' in result['notice']


async def test_malformed_workbook_refuses_without_returning_parser_details(output_state):
    state = output_state
    conversation, run = await _run(state)
    item = await _create(state, conversation, run, filename='budget.xlsx', format='xlsx', sheets=[{'name': 'Budget', 'rows': [['Cost', 40]]}])
    from test_generated_files import httpx_request

    from audrey.routes import files as uploads
    row = await state.db.get_upload(item['file_id'])
    uploads._source_path(httpx_request(state.app), row).write_bytes(b'PRIVATE_PARSED_CONTENT')
    response = await state.client.get(f"/api/files/{item['file_id']}/preview")
    assert response.status_code == 422
    assert response.json()['detail'] == 'Preview is unavailable. You can still download the file.'
    assert 'PRIVATE_PARSED' not in response.text


def test_spreadsheet_columns_and_formula_text_remain_literal():
    source = io.BytesIO()
    book = Workbook()
    book.active.append(['=SUM(1,2)', True, None, 'text'])
    book.save(source)
    source.seek(0)
    result = workbook_preview(source)
    assert result['sheets'][0]['rows'][0] == ['=SUM(1,2)', True, None, 'text']
    assert result['truncated'] is False


def test_bounded_archive_rejects_inflation_before_openpyxl():
    source = io.BytesIO()
    with zipfile.ZipFile(source, 'w', compression=zipfile.ZIP_DEFLATED) as package:
        package.writestr('huge.xml', b'0' * (17 * 1024 * 1024))
    source.seek(0)
    with pytest.raises(FilePreviewError, match='large'):
        workbook_preview(source)


def test_document_paragraph_count_is_bounded(tmp_path):
    from audrey.file_previews import MAX_PARAGRAPHS, preview_generated_file
    source = tmp_path / 'notes.docx'
    source.write_bytes(b'generated source')
    text = tmp_path / 'notes.ocr.txt'
    text.write_text('\n\n'.join(['Paragraph'] * (MAX_PARAGRAPHS + 1)))
    result = preview_generated_file(source, format='docx', text_path=text)
    assert len(result['paragraphs']) == MAX_PARAGRAPHS
    assert result['truncated']


@pytest.mark.parametrize('format', ['csv', 'xlsx'])
async def test_table_pages_have_no_missing_rows_and_repeat_ownership_checks(output_state, format):
    state = output_state
    conversation, run = await _run(state)
    rows = [[index, f'row-{index}'] for index in range(405)]
    fields = {'rows': rows} if format == 'csv' else {'sheets': [
        {'name': 'First', 'rows': rows}, {'name': 'Second', 'rows': [['second sheet']]},
    ]}
    item = await _create(state, conversation, run, filename=f'paged.{format}', format=format, **fields)
    path = f"/api/files/{item['file_id']}/preview"
    seen = []
    for offset, expected_next in [(0, 200), (200, 400), (400, None)]:
        response = await state.client.get(path, params={'sheet': 0, 'row_offset': offset})
        assert response.status_code == 200, response.text
        page = response.json()
        table = page['sheets'][0]
        assert table['sheet_index'] == 0 and table['row_offset'] == offset
        assert table['next_row_offset'] == expected_next
        assert table['total_rows'] == (405 if offset == 400 else None)
        assert page['content_truncated'] is False and table['content_truncated'] is False
        assert response.headers['cache-control'] == 'private, no-store'
        seen.extend(table['rows'])
    expected = [[str(cell) for cell in row] for row in rows] if format == 'csv' else rows
    assert seen == expected
    if format == 'xlsx':
        result = (await state.client.get(path, params={'sheet': 1})).json()
        assert result['sheet_count'] == 2 and len(result['sheets']) == 1
        assert result['sheets'][0]['sheet_index'] == 1
        assert result['sheets'][0]['rows'] == [['second sheet']]
    state.current[0] = state.other
    assert (await state.client.get(path, params={'sheet': 0, 'row_offset': 200})).status_code == 404
    state.current[0] = state.owner
    assert (await state.client.delete(f"/api/files/{item['file_id']}")).status_code == 200
    assert (await state.client.get(path, params={'sheet': 0, 'row_offset': 200})).status_code == 404


@pytest.mark.parametrize('format', ['txt', 'md', 'docx', 'json'])
async def test_text_pages_join_exactly_including_unicode_and_paragraph_boundary(output_state, format):
    state = output_state
    conversation, run = await _run(state)
    content = ('Paragraph \u2022 with text\n\n' * 6000)[:150_000]
    fields = {'data': {'note': content}} if format == 'json' else {'content': content}
    item = await _create(state, conversation, run, filename=f'long.{format}', format=format, **fields)
    assert 'error' not in item, item
    path = f"/api/files/{item['file_id']}/preview"
    chunks, offset = [], 0
    while True:
        response = await state.client.get(path, params={'text_offset': offset})
        assert response.status_code == 200, response.text
        result = response.json()
        assert result['offset'] == offset
        assert len(result['text']) <= 100_000
        assert result['content_truncated'] is False
        if format == 'docx':
            assert len(result['paragraphs']) <= 1000
        chunks.append(result['text'])
        next_offset = result['next_offset']
        if next_offset is None:
            break
        assert next_offset == offset + len(result['text'])
        assert next_offset > offset
        offset = next_offset
    if format == 'json':
        import json
        assert json.loads(''.join(chunks)) == fields['data']
    else:
        assert ''.join(chunks) == content
    assert sum(map(len, chunks)) == result['total_chars']


@pytest.mark.parametrize('params', [
    {'sheet': -1}, {'sheet': 8}, {'sheet': 'oops'},
    {'row_offset': -1}, {'row_offset': 10001},
    {'text_offset': -1}, {'text_offset': 8388609},
])
async def test_preview_positions_are_bounded_by_route(output_state, params):
    response = await output_state.client.get('/api/files/not-a-file/preview', params=params)
    assert response.status_code == 422


def test_table_continuation_does_not_hide_irrecoverable_cell_or_column_omissions():
    source = io.BytesIO()
    book = Workbook()
    book.active.append(['text' * 3000] + [None] * 32 + ['hidden final column'])
    for index in range(205):
        book.active.append([index])
    book.save(source)
    source.seek(0)
    first = workbook_preview(source)
    sheet = first['sheets'][0]
    assert sheet['next_row_offset'] == MAX_ROWS
    assert sheet['content_truncated'] and first['content_truncated']
    assert len(sheet['rows'][0][0]) == 8000


def test_invalid_sheet_refuses_instead_of_showing_another_sheet():
    source = io.BytesIO()
    book = Workbook()
    book.save(source)
    source.seek(0)
    with pytest.raises(FilePreviewError, match='Worksheet'):
        workbook_preview(source, sheet=1)

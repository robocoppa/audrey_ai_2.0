"""Generated artifacts preserve literal content and enforce bounded schemas."""
from __future__ import annotations

import csv
import hashlib
import io
import json
import zipfile
from datetime import date

import pytest
from docx import Document
from openpyxl import load_workbook
from pydantic import ValidationError

from audrey.file_rendering import (
    MAX_INPUT_BYTES,
    CreateFileRequest,
    FileSheet,
    render_file,
)


@pytest.mark.parametrize("format", ["txt", "md"])
def test_plain_text_and_markdown_are_utf8_without_processing(format):
    content = "# Notes\n\nKeep `=1+1` and <script>literal text</script>.\nRésumé"
    artifact = render_file(CreateFileRequest(filename=f"notes.{format}", format=format,
                                             content=content))
    assert artifact.data.decode() == content == artifact.text
    assert artifact.sha256 == hashlib.sha256(artifact.data).hexdigest()
    assert artifact.mime_type.startswith("text/")


def test_json_file_keeps_json_types_and_accepts_explicit_null():
    value = {"string": "blue canoe", "values": [1, 1.5, True, None], "nested": {"a": []}}
    artifact = render_file(CreateFileRequest(filename="report.json", format="json", data=value))
    assert json.loads(artifact.data) == value
    assert artifact.text == artifact.data.decode()
    assert artifact.mime_type == "application/json"
    assert render_file(CreateFileRequest(filename="null.json", format="json", data=None)).data == b"null\n"


def test_csv_handles_commas_newlines_unicode_and_neutralizes_spreadsheet_formulas():
    original = [
        ["Item", "Value", "Unicode"],
        ["comma, quote\"\nnext line", 12.5, "café"],
        ["=HYPERLINK(\"https://outside.invalid\",\"click\")", "+SUM(1,2)", "-20"],
        ["@SUM(1,2)", " \t=1+1", "\tplain"],
        ["\rplain", "\nplain", -20],
        [True, False, None],
    ]
    artifact = render_file(CreateFileRequest(filename="report.csv", format="csv", rows=original))
    rows = list(csv.reader(io.StringIO(artifact.data.decode("utf-8-sig"))))
    assert rows[1] == [original[1][0], "12.5", "café"]
    assert rows[2] == ["'" + value for value in original[2]]
    assert rows[3] == ["'" + value for value in original[3]]
    assert rows[4] == ["'\rplain", "'\nplain", "-20"]
    assert rows[5] == ["True", "False", ""]
    # A local viewer uses the original values, not CSV's protective apostrophes.
    assert list(csv.reader(io.StringIO(artifact.text)))[2] == original[2]


def _assert_no_external_or_executable_zip_parts(data: bytes):
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        assert all("vbaProject" not in name and "externalLinks" not in name
                   for name in archive.namelist())
        for name in archive.namelist():
            if name.endswith(".rels"):
                assert b'TargetMode="External"' not in archive.read(name)


def test_docx_is_a_real_document_with_plain_paragraphs_and_no_external_resources():
    content = "Overview\n\n=HYPERLINK(\"https://outside.invalid\",\"click\")\nRésumé"
    artifact = render_file(CreateFileRequest(filename="notes.docx", format="docx", content=content))
    document = Document(io.BytesIO(artifact.data))
    assert [paragraph.text for paragraph in document.paragraphs] == content.split("\n")
    assert artifact.text == content
    assert artifact.mime_type.endswith("wordprocessingml.document")
    _assert_no_external_or_executable_zip_parts(artifact.data)


def test_xlsx_is_a_real_workbook_with_literal_formulas_and_exact_large_integers():
    rows = [["Label", "Value"], ["formula", "=1+1"], ["link", "https://outside.invalid"],
            ["number", 12.5], ["boolean", True], ["big integer", 12345678901234567890],
            ["error literal", "#REF!"], ["negative string", "-20"]]
    artifact = render_file(CreateFileRequest(filename="report.xlsx", format="xlsx",
        sheets=[FileSheet(name="Summary", rows=rows), FileSheet(name="Details", rows=[[None]])]))
    workbook = load_workbook(io.BytesIO(artifact.data), data_only=False)
    assert workbook.sheetnames == ["Summary", "Details"]
    assert workbook["Summary"]["B2"].value == "=1+1"
    assert workbook["Summary"]["B2"].data_type == "s"
    assert workbook["Summary"]["B3"].hyperlink is None
    assert workbook["Summary"]["B4"].value == 12.5
    assert workbook["Summary"]["B5"].value is True
    assert workbook["Summary"]["B6"].value == "12345678901234567890"
    assert workbook["Summary"]["B7"].data_type == "s"
    assert not any(cell.data_type == "f" for sheet in workbook for row in sheet for cell in row)
    workbook.close()
    assert "Summary" in artifact.text and "=1+1" in artifact.text
    _assert_no_external_or_executable_zip_parts(artifact.data)


@pytest.mark.parametrize("filename", [
    "../notes.txt", "dir/notes.txt", "dir\\notes.txt", "/notes.txt", "C:notes.txt",
    ".hidden.txt", " notes.txt", "notes.txt ", "notes.txt.", "CON.txt", "lpt1.txt",
    "notes\n.txt", "notes\x00.txt", "notes\u202e.txt", "*.txt", "?.txt", ".txt",
    "notes.md", "a" * 177 + ".txt",
])
def test_unsafe_or_mismatching_filenames_are_rejected(filename):
    with pytest.raises(ValidationError):
        CreateFileRequest(filename=filename, format="txt", content="text")


def test_uppercase_suffix_is_accepted_and_no_filesystem_paths_are_allowed(tmp_path, monkeypatch):
    def no_open(*args, **kwargs):
        pytest.fail("Rendering must not read or write filesystem paths")
    monkeypatch.setattr("builtins.open", no_open)
    artifact = render_file(CreateFileRequest(filename="Notes.TXT", format="txt", content="hello"))
    assert artifact.data == b"hello"
    assert list(tmp_path.iterdir()) == []


@pytest.mark.parametrize("changes", [
    {}, {"content": None}, {"content": "text", "rows": [["mixed"]]},
    {"content": "text", "data": None}, {"content": "text", "path": "/etc/passwd"},
    {"content": "text", "url": "https://outside.invalid"},
    {"content": "text", "content_template": "expression"},
])
def test_content_fields_are_required_exclusive_and_extra_fields_forbidden(changes):
    with pytest.raises(ValidationError):
        CreateFileRequest(filename="notes.txt", format="txt", **changes)


@pytest.mark.parametrize("format,values", [
    ("json", {}), ("json", {"data": {"x": 1}, "content": None}),
    ("csv", {"rows": []}), ("csv", {"rows": [[{"nested": "not a cell"}]]}),
    ("xlsx", {"sheets": []}), ("xlsx", {"sheets": [{"name": "Summary", "rows": [[1]], "formula": "=1"}]}),
    ("pdf", {"content": "not supported"}),
])
def test_invalid_format_fields_and_cell_shapes_are_rejected(format, values):
    with pytest.raises(ValidationError):
        CreateFileRequest(filename=f"report.{format}", format=format, **values)


@pytest.mark.parametrize("name", ["", " ", " Summary", "Summary ", "Sheet/1", "Sheet\\1",
                                  "[1]", "a:b", "a?b", "a*b", "'Sheet", "Sheet'",
                                  "\tSheet", "Sheet\u202e", "x" * 32])
def test_worksheet_names_are_safe(name):
    with pytest.raises(ValidationError):
        CreateFileRequest(filename="report.xlsx", format="xlsx",
                          sheets=[{"name": name, "rows": [[1]]}])


def test_duplicate_worksheet_names_and_aggregate_row_limit_are_rejected():
    with pytest.raises(ValidationError, match="unique"):
        CreateFileRequest(filename="report.xlsx", format="xlsx", sheets=[
            {"name": "Summary", "rows": [[1]]}, {"name": "summary", "rows": [[2]]}])
    with pytest.raises(ValidationError, match="row limit"):
        CreateFileRequest(filename="report.xlsx", format="xlsx", sheets=[
            {"name": "First", "rows": [[1]] * 501}, {"name": "Second", "rows": [[1]] * 500}])
    with pytest.raises(ValidationError):
        CreateFileRequest(filename="report.xlsx", format="xlsx", sheets=[
            {"name": f"Sheet {number}", "rows": [[1]]} for number in range(9)])


@pytest.mark.parametrize("format,values", [
    ("txt", {"content": "a" * 200001}),
    ("txt", {"content": "\U0001f6f6" * 100000}),
    ("txt", {"content": "null\x00character"}),
    ("csv", {"rows": [["a" * 32768]]}),
    ("csv", {"rows": [[1] * 33]}),
    ("csv", {"rows": [[1]] * 1001}),
    ("csv", {"rows": [[float("nan")]]}),
    ("xlsx", {"sheets": [{"name": "Summary", "rows": [[float("inf")]]}]}),
    ("docx", {"content": "surrogate\ud800"}),
])
def test_text_and_cell_bounds(format, values):
    with pytest.raises(ValidationError):
        CreateFileRequest(filename=f"notes.{format}", format=format, **values)


@pytest.mark.parametrize("data", [float("nan"), {"key": float("inf")}, {1: "wrong key"},
                                 {"date": date(2025, 1, 1)}, {"bytes": b"not JSON"},
                                 {"set": {1, 2}}, "surrogate\ud800"])
def test_json_is_strictly_json_with_finite_numbers(data):
    with pytest.raises(ValidationError):
        CreateFileRequest(filename="report.json", format="json", data=data)


def test_json_depth_and_aggregate_input_are_bounded():
    nested = None
    for _ in range(30):
        nested = [nested]
    for data in (nested, [0] * 32001, {"large": "x" * MAX_INPUT_BYTES}):
        with pytest.raises(ValidationError):
            CreateFileRequest(filename="report.json", format="json", data=data)
    recursive = []
    recursive.append(recursive)
    with pytest.raises(ValidationError, match="nesting"):
        CreateFileRequest(filename="report.json", format="json", data=recursive)


def test_renderer_revalidates_models_constructed_without_validation_and_caps_output(monkeypatch):
    request = CreateFileRequest.model_construct(filename="notes.txt", format="txt", content="\x00")
    with pytest.raises(ValidationError):
        render_file(request)
    monkeypatch.setattr("audrey.file_rendering.MAX_OUTPUT_BYTES", 10)
    with pytest.raises(ValueError, match="output limit"):
        render_file(CreateFileRequest(filename="notes.txt", format="txt", content="long text here"))


def test_schema_is_compact_and_only_one_nonrecursive_definition_is_needed():
    schema = CreateFileRequest.model_json_schema()
    assert schema["additionalProperties"] is False
    assert set(schema["$defs"]) == {"FileSheet"}
    assert schema["$defs"]["FileSheet"]["additionalProperties"] is False
    assert len(json.dumps(schema)) < 6000
    assert schema["properties"]["format"]["enum"] == ["txt", "md", "csv", "json", "docx", "xlsx"]


def test_markdown_docx_saves_native_headings_runs_lists_table_and_code():
    content = (
        '# Review notes\n\n## Findings\nUse **verified** facts and *clear* caveats with `literal` code.\n'
        '- First observation\n- Second observation\n1. Read the source\n2. Check the figure\n'
        '> Keep this qualification.\n\n| Item | Amount |\n| --- | --- |\n| Wages | **50000** |\n'
        '\n```python\n=1+1\nprint("literal")\n```'
    )
    artifact = render_file(CreateFileRequest(filename='review.docx', format='docx',
                                             content_format='markdown', content=content))
    document = Document(io.BytesIO(artifact.data))
    assert document.paragraphs[0].style.name == 'Heading 1'
    assert document.paragraphs[0].text == 'Review notes'
    assert document.paragraphs[1].style.name == 'Heading 2'
    rich = document.paragraphs[2]
    assert any(run.text == 'verified' and run.bold for run in rich.runs)
    assert any(run.text == 'clear' and run.italic for run in rich.runs)
    assert any(run.text == 'literal' and run.style.name == 'InlineCode' for run in rich.runs)
    assert [paragraph.style.name for paragraph in document.paragraphs[3:7]] == [
        'List Bullet', 'List Bullet', 'List Number', 'List Number']
    assert document.paragraphs[7].style.name == 'Quote'
    assert document.paragraphs[-1].style.name == 'CodeBlock'
    assert document.paragraphs[-1].text == '=1+1\nprint("literal")'
    assert [[cell.text for cell in row.cells] for row in document.tables[0].rows] == [
        ['Item', 'Amount'], ['Wages', '50000']]
    assert all(run.bold for cell in document.tables[0].rows[0].cells for run in cell.paragraphs[0].runs)
    assert 'Item\tAmount\nWages\t50000' in artifact.text
    assert 'Use verified facts and clear caveats with literal code.' in artifact.text
    assert artifact.text.startswith('Review notes\n\nFindings\n\n')
    _assert_no_external_or_executable_zip_parts(artifact.data)


def test_docx_markdown_is_opt_in_and_keeps_unsupported_markup_literal():
    content = '# Heading\n![image](https://outside.invalid/a.png)\n<script>literal()</script>'
    plain = render_file(CreateFileRequest(filename='plain.docx', format='docx', content=content))
    assert Document(io.BytesIO(plain.data)).paragraphs[0].text == '# Heading'
    assert plain.text == content
    rich = render_file(CreateFileRequest(filename='rich.docx', format='docx',
                                         content_format='markdown', content=content))
    document = Document(io.BytesIO(rich.data))
    assert document.paragraphs[0].text == 'Heading'
    assert document.paragraphs[1].text == '![image](https://outside.invalid/a.png)'
    assert document.paragraphs[2].text == '<script>literal()</script>'
    _assert_no_external_or_executable_zip_parts(rich.data)
    unclosed = render_file(CreateFileRequest(filename='fence.docx', format='docx',
        content_format='markdown', content='```\nliteral text'))
    assert unclosed.text == '```\nliteral text'


def test_docx_markdown_table_and_block_work_are_bounded():
    with pytest.raises(ValueError, match='column limit'):
        render_file(CreateFileRequest(filename='wide.docx', format='docx', content_format='markdown',
            content='|'.join(['item'] * 33) + '\n' + '|'.join(['---'] * 33)))
    with pytest.raises(ValueError, match='block limit'):
        render_file(CreateFileRequest(filename='many.docx', format='docx', content_format='markdown',
                                      content='line\n' * 2001))


def test_xlsx_styles_actual_headers_widths_formats_without_changing_cells():
    rows = [['Item', 'Count', 'Amount', 'Fraction', 'Literal'],
            ['Wages', 1234, 50000.125, 0.125, '=SUM(1,2)'],
            ['Other', 10, -12.50, 1, '#REF!']]
    artifact = render_file(CreateFileRequest(filename='styled.xlsx', format='xlsx', sheets=[{
        'name': 'Worksheet', 'rows': rows, 'header_row': True,
        'column_formats': ['text', 'integer', 'currency', 'percent', 'general'],
        'column_widths': [25, 10, 20],
    }]))
    book = load_workbook(io.BytesIO(artifact.data), data_only=False)
    sheet = book['Worksheet']
    assert sheet.freeze_panes == 'A2'
    assert sheet.auto_filter.ref == 'A1:E3'
    assert sheet['A1'].font.bold and sheet['A1'].font.color.rgb == '00FFFFFF'
    assert sheet['A1'].fill.fgColor.rgb == '0016324F'
    assert sheet['A2'].fill.fgColor.rgb == '00F2F6FA'
    assert [sheet[f'{column}2'].number_format for column in 'ABCDE'] == [
        '@', '#,##0', '$#,##0.00', '0.00%', 'General']
    assert [sheet[f'{column}2'].value for column in 'ABCDE'] == rows[1]
    assert sheet['E2'].data_type == 's' and sheet['E2'].hyperlink is None
    assert sheet.column_dimensions['A'].width == 25
    assert sheet.column_dimensions['B'].width == 10
    assert sheet.column_dimensions['E'].width >= 10
    assert all(cell.alignment.horizontal == 'left' for row in sheet for cell in row)
    book.close()
    _assert_no_external_or_executable_zip_parts(artifact.data)


def test_default_xlsx_style_does_not_guess_header_or_number_formats():
    artifact = render_file(CreateFileRequest(filename='values.xlsx', format='xlsx', sheets=[{
        'name': 'Values', 'rows': [[1, 0.125], [2, 0.25]],
    }]))
    book = load_workbook(io.BytesIO(artifact.data))
    sheet = book['Values']
    assert sheet.freeze_panes is None and sheet.auto_filter.ref is None
    assert sheet['A1'].font.bold is False
    assert sheet['B1'].value == 0.125 and sheet['B1'].number_format == 'General'
    assert all(cell.alignment.horizontal == 'left' for row in sheet for cell in row)
    book.close()


@pytest.mark.parametrize('options', [
    {'column_formats': ['formula']}, {'column_formats': ['currency', 'integer']},
    {'column_widths': [7]}, {'column_widths': [61]}, {'column_widths': [True]},
    {'column_widths': [float('nan')]}, {'column_widths': [10, 20]}, {'header_row': 'true'},
])
def test_worksheet_formatting_options_are_safe_bounded_and_typed(options):
    with pytest.raises(ValidationError):
        CreateFileRequest(filename='report.xlsx', format='xlsx',
                          sheets=[{'name': 'Sheet', 'rows': [[1]], **options}])


@pytest.mark.parametrize('format', ['txt', 'md', 'csv', 'json', 'xlsx'])
def test_content_format_option_is_docx_only(format):
    field = {'txt': {'content': 'text'}, 'md': {'content': '# text'}, 'csv': {'rows': [[1]]},
             'json': {'data': 1}, 'xlsx': {'sheets': [{'name': 'Sheet', 'rows': [[1]]}]}}[format]
    with pytest.raises(ValidationError, match='only for DOCX'):
        CreateFileRequest(filename=f'report.{format}', format=format, content_format='plain', **field)


def test_docx_markdown_preserves_heading_hashes_and_explicit_list_numbers():
    artifact = render_file(CreateFileRequest(filename='numbered.docx', format='docx',
        content_format='markdown', content='# C#\n4. Fourth item\n7. Seventh item\n\n1. New list'))
    document = Document(io.BytesIO(artifact.data))
    assert document.paragraphs[0].text == 'C#'
    numbering = document.part.numbering_part.element
    starts = []
    for paragraph in document.paragraphs[1:]:
        identity = paragraph._p.pPr.numPr.numId.val
        instance = numbering.num_having_numId(identity)
        starts.append(instance.lvlOverride_lst[0].startOverride.val)
    assert starts == [4, 7, 1]

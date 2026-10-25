"""Workbook exports preserve saved authority, ownership and bounded lifecycle."""
from __future__ import annotations

import asyncio
import copy
import json
import threading
import zipfile

import pytest
from openpyxl import load_workbook

from audrey.tax_exports import TaxExportError, TaxReviewExports


def _case(**changes):
    result = {
        "id": "case1", "conversation_id": "con1", "project_id": "project1",
        "tax_year": 2025, "jurisdiction": "US-federal", "revision": 8,
        "reference_release": "release1", "goal": "Compare selected sources",
        "facts": [{"id": "fact1", "label": "Wages", "value": "50000.00",
                   "status": "confirmed", "case_tax_year": 2025,
                   "confirmed_tax_year": 2025, "history": [{"value": "49000.00"}]}],
        "evidence": [{"id": "evidence1", "claim": "Source condition", "document_id": "a" * 64,
                      "revision": "b" * 64, "reference_release": "release1", "page": 1,
                      "passage": "The stated amount is 15750 for 2025.", "tax_year": 2025,
                      "tax_year_status": "proposed", "status": "needs_recheck",
                      "conflicts_with": ["evidence2"], "recheck_for_reference_release": "release2"}],
        "calculations": [{"id": "calc1", "label": "Conditional subtraction",
                          "inputs": {"wages": "50000.00", "deduction": "15750.00"},
                          "expression": "wages - deduction", "rounding": "none",
                          "result": "34250.00", "fact_ids": ["fact1"],
                          "evidence_ids": ["evidence1"], "case_tax_year": 2025,
                          "status": "recorded"}],
        "return_items": [{"id": "item1", "label": "Conditional amount", "status": "proposed",
                          "fact_ids": ["fact1"], "evidence_ids": ["evidence1"]}],
        "missing_items": [{"id": "missing1", "label": "Source conflict",
                           "question": "Which rule applies?", "status": "open"}],
        "review": {"review_revision": 8, "summary": "Useful partial review",
                   "saved_at": "2026-10-24T00:00:00+00:00", "status": "limited_evidence",
                   "stale": False, "limitations": ["A source conflict is unresolved."],
                   "next_question": "", "findings": [{"label": "Conditional result",
                        "detail": "Supplied wages minus the source amount is 34250.00.",
                        "fact_ids": ["fact1"], "evidence_ids": ["evidence1"],
                        "calculation_ids": ["calc1"], "return_item_ids": ["item1"]}]},
        "export_sources": {"evidence1": {"title": "Instructions for a form",
                           "source_url": "https://www.irs.gov/pub/irs-pdf/i8911.pdf",
                           "revision": "December 2025"}},
    }
    return {**result, **changes}


def _cells(workbook, sheet):
    return [cell for row in workbook[sheet].iter_rows() for cell in row if cell.value is not None]


async def test_prepare_is_lazy_deterministic_and_keeps_private_data_in_private_files(tmp_path):
    exports = TaxReviewExports(tmp_path / "exports")
    case = _case()
    before = copy.deepcopy(case)
    first = await exports.prepare(user_id="owner", case=case)
    second = await exports.prepare(user_id="owner", case=case)
    assert first == second and case == before
    assert first["case_revision"] == first["review_revision"] == 8
    assert first["download_url"] == f"/api/conversations/con1/tax-review/exports/{first['id']}/download"
    assert "case-r8-review-r8.xlsx" in first["filename"] and first["expires_at"]
    snapshots = list(exports.root.glob("*/*/*.json"))
    assert len(snapshots) == 1 and not list(exports.root.rglob("*.xlsx"))
    assert snapshots[0].stat().st_mode & 0o777 == 0o600
    assert snapshots[0].parent.stat().st_mode & 0o777 == 0o700
    assert not any("owner" in str(path.relative_to(exports.root)) for path in snapshots)


async def test_workbook_contains_complete_saved_values_scope_conflicts_and_original_citations(tmp_path):
    exports = TaxReviewExports(tmp_path)
    case = _case()
    reference = await exports.prepare(user_id="owner", case=case)
    artifact = await exports.download(user_id="owner", case=case, export_id=reference["id"])
    assert artifact.case_revision == 8 and artifact.review_revision == 8
    assert artifact.filename == reference["filename"]
    assert artifact.path.stat().st_mode & 0o777 == 0o600
    workbook = load_workbook(artifact.path)
    assert workbook.sheetnames == ["Overview", "Facts", "Evidence", "Calculations", "Unresolved items"]
    overview = [cell.value for cell in _cells(workbook, "Overview")]
    assert any("does not determine eligibility" in str(value) for value in overview)
    assert "Useful partial review" in overview
    assert "49000.00" in str([cell.value for cell in _cells(workbook, "Facts")])
    evidence = [cell.value for cell in _cells(workbook, "Evidence")]
    assert "December 2025" in evidence and case["evidence"][0]["passage"] in evidence
    assert case["export_sources"]["evidence1"]["source_url"] in evidence
    assert "evidence2" in str(evidence) and "proposed" in evidence
    assert "34250.00" in [cell.value for cell in _cells(workbook, "Calculations")]
    unresolved = [cell.value for cell in _cells(workbook, "Unresolved items")]
    assert "Which rule applies?" in unresolved and "needs_recheck" in unresolved
    workbook.close()


async def test_strings_never_become_formulas_links_or_lossy_decimal_numbers(tmp_path):
    case = _case(goal='=HYPERLINK("https://outside.invalid", "click")')
    case["facts"][0]["value"] = "+SUM(1,2)"
    case["calculations"][0]["expression"] = "=1+1"
    case["calculations"][0]["result"] = "123456789012345678901234567890.000001"
    exports = TaxReviewExports(tmp_path)
    reference = await exports.prepare(user_id="owner", case=case)
    artifact = await exports.download(user_id="owner", case=case, export_id=reference["id"])
    workbook = load_workbook(artifact.path, data_only=False)
    all_cells = [cell for name in workbook.sheetnames for cell in _cells(workbook, name)]
    assert all(cell.data_type == "s" and cell.hyperlink is None for cell in all_cells)
    assert case["goal"] in [cell.value for cell in all_cells]
    assert case["calculations"][0]["result"] in [cell.value for cell in all_cells]
    workbook.close()
    with zipfile.ZipFile(artifact.path) as archive:
        assert not any("externalLinks" in name or "vbaProject" in name for name in archive.namelist())


async def test_revision_link_survives_later_case_changes_without_changing_its_contents(tmp_path):
    exports = TaxReviewExports(tmp_path)
    original = _case()
    reference = await exports.prepare(user_id="owner", case=original)
    changed = copy.deepcopy(original)
    changed.update(revision=9, tax_year=2026)
    changed["review"].update(stale=True, summary="New summary")
    original["facts"][0]["value"] = "MUTATED AFTER PREPARE"
    artifact = await exports.download(user_id="owner", case=changed, export_id=reference["id"])
    workbook = load_workbook(artifact.path)
    assert "50000.00" in [cell.value for cell in _cells(workbook, "Facts")]
    assert "Useful partial review" in [cell.value for cell in _cells(workbook, "Overview")]
    assert "2025" in [cell.value for cell in _cells(workbook, "Overview")]
    workbook.close()


async def test_concurrent_downloads_render_once_and_cache_survives_service_restart(tmp_path, monkeypatch):
    exports = TaxReviewExports(tmp_path)
    case = _case()
    reference = await exports.prepare(user_id="owner", case=case)
    render = exports._render
    count = 0

    def measured(*args):
        nonlocal count
        count += 1
        render(*args)

    monkeypatch.setattr(exports, "_render", measured)
    artifacts = await asyncio.gather(*[exports.download(user_id="owner", case=case, export_id=reference["id"]) for _ in range(4)])
    assert count == 1 and len({item.path for item in artifacts}) == 1
    restarted = TaxReviewExports(tmp_path)
    monkeypatch.setattr(restarted, "_render", lambda *args: pytest.fail("cached file was regenerated"))
    again = await restarted.download(user_id="owner", case=case, export_id=reference["id"])
    assert again.path == artifacts[0].path


@pytest.mark.parametrize("change", ["owner", "case", "conversation", "path"])
async def test_download_scope_cannot_cross_owner_case_conversation_or_directory(tmp_path, change):
    exports = TaxReviewExports(tmp_path)
    case = _case()
    reference = await exports.prepare(user_id="owner", case=case)
    user_id, export_id = "owner", reference["id"]
    if change == "owner":
        user_id = "another"
    elif change == "case":
        case["id"] = "another"
    elif change == "conversation":
        case["conversation_id"] = "another"
    else:
        export_id = "../../outside"
    with pytest.raises(TaxExportError) as error:
        await exports.download(user_id=user_id, case=case, export_id=export_id)
    assert error.value.status_code == 404
    assert not list(exports.root.rglob("*.xlsx"))


async def test_snapshot_tampering_is_refused_without_exposing_contents(tmp_path):
    exports = TaxReviewExports(tmp_path)
    case = _case()
    reference = await exports.prepare(user_id="owner", case=case)
    path = next(exports.root.glob("*/*/*.json"))
    snapshot = json.loads(path.read_bytes())
    snapshot["owner"] = "another"
    path.write_text(json.dumps(snapshot))
    with pytest.raises(TaxExportError) as error:
        await exports.download(user_id="owner", case=case, export_id=reference["id"])
    assert error.value.status_code == 404 and "another" not in str(error.value)


async def test_expiry_removes_snapshot_and_output(tmp_path):
    exports = TaxReviewExports(tmp_path)
    case = _case()
    reference = await exports.prepare(user_id="owner", case=case)
    artifact = await exports.download(user_id="owner", case=case, export_id=reference["id"])
    path = next(exports.root.glob("*/*/*.json"))
    snapshot = json.loads(path.read_bytes())
    snapshot["expires_at"] = 1
    path.write_text(json.dumps(snapshot))
    with pytest.raises(TaxExportError) as error:
        await exports.download(user_id="owner", case=case, export_id=reference["id"])
    assert error.value.status_code == 410 and not path.exists() and not artifact.path.exists()


async def test_retention_bounds_revisions_and_global_snapshots(tmp_path):
    exports = TaxReviewExports(tmp_path, max_case_exports=2, max_snapshots=3)
    references = []
    for revision in range(1, 5):
        references.append(await exports.prepare(user_id="owner", case=_case(revision=revision)))
    assert len(list(exports.root.rglob("*.json"))) == 2
    with pytest.raises(TaxExportError) as error:
        await exports.download(user_id="owner", case=_case(), export_id=references[0]["id"])
    assert error.value.status_code == 404
    await exports.prepare(user_id="second", case=_case())
    await exports.prepare(user_id="third", case=_case())
    assert len(list(exports.root.rglob("*.json"))) == 3


@pytest.mark.parametrize("limits", [{"max_snapshot_bytes": 100}, {"max_cache_bytes": 100}])
async def test_snapshot_and_cache_limits_leave_the_canonical_case_unchanged(tmp_path, limits):
    exports = TaxReviewExports(tmp_path, **limits)
    case = _case()
    before = copy.deepcopy(case)
    with pytest.raises(TaxExportError) as error:
        await exports.prepare(user_id="owner", case=case)
    assert error.value.status_code == 422 and case == before
    assert not list(exports.root.rglob("*.json"))


@pytest.mark.parametrize("limits", [{"max_rows": 2}, {"max_workbook_bytes": 100}])
async def test_render_limits_keep_snapshot_retryable_and_never_publish_partial_file(tmp_path, limits):
    exports = TaxReviewExports(tmp_path, **limits)
    case = _case()
    reference = await exports.prepare(user_id="owner", case=case)
    with pytest.raises(TaxExportError) as error:
        await exports.download(user_id="owner", case=case, export_id=reference["id"])
    assert error.value.status_code == 422
    assert len(list(exports.root.rglob("*.json"))) == 1
    assert not list(exports.root.rglob("*.xlsx")) and not list(exports.root.rglob("*.tmp"))


async def test_failed_render_is_public_safe_and_can_be_retried(tmp_path, monkeypatch):
    exports = TaxReviewExports(tmp_path)
    case = _case()
    reference = await exports.prepare(user_id="owner", case=case)
    original = exports._render

    def fail(*args):
        raise OSError("private document contents must not reach the client")

    monkeypatch.setattr(exports, "_render", fail)
    with pytest.raises(TaxExportError) as error:
        await exports.download(user_id="owner", case=case, export_id=reference["id"])
    assert error.value.status_code == 503 and "private document contents" not in str(error.value)
    monkeypatch.setattr(exports, "_render", original)
    assert (await exports.download(user_id="owner", case=case, export_id=reference["id"])).path.exists()


async def test_case_deletion_invalidates_inflight_render_and_prevents_republication(tmp_path, monkeypatch):
    exports = TaxReviewExports(tmp_path)
    case = _case()
    reference = await exports.prepare(user_id="owner", case=case)
    original = exports._render
    started, finish = threading.Event(), threading.Event()

    def blocked(*args):
        started.set()
        assert finish.wait(5)
        original(*args)

    monkeypatch.setattr(exports, "_render", blocked)
    rendering = asyncio.create_task(exports.download(user_id="owner", case=case, export_id=reference["id"]))
    assert await asyncio.to_thread(started.wait, 5)
    deleting = asyncio.create_task(exports.delete_case(user_id="owner", case_id=case["id"]))
    await asyncio.sleep(0)
    finish.set()
    with pytest.raises(TaxExportError) as error:
        await rendering
    await deleting
    assert error.value.status_code == 404
    assert not list(exports.root.rglob("*.json")) and not list(exports.root.rglob("*.xlsx"))
    with pytest.raises(TaxExportError):
        await exports.prepare(user_id="owner", case=case)
    # A rejected canonical deletion can safely reopen exports; old queued work
    # stays invalidated and a fresh request captures canonical data again.
    await exports.allow_case(user_id="owner", case_id=case["id"])
    assert await exports.prepare(user_id="owner", case=case)


async def test_cancelled_render_waits_for_worker_then_removes_temporary_output(tmp_path, monkeypatch):
    exports = TaxReviewExports(tmp_path)
    case = _case()
    reference = await exports.prepare(user_id="owner", case=case)
    original = exports._render
    started, finish = threading.Event(), threading.Event()

    def blocked(*args):
        started.set()
        assert finish.wait(5)
        original(*args)

    monkeypatch.setattr(exports, "_render", blocked)
    rendering = asyncio.create_task(exports.download(user_id="owner", case=case, export_id=reference["id"]))
    assert await asyncio.to_thread(started.wait, 5)
    rendering.cancel()
    await asyncio.sleep(0)
    assert not rendering.done()
    finish.set()
    with pytest.raises(asyncio.CancelledError):
        await rendering
    assert not list(exports.root.rglob("*.xlsx")) and not list(exports.root.rglob("*.tmp"))


async def test_owner_purge_removes_every_old_export_but_allows_new_cases(tmp_path):
    exports = TaxReviewExports(tmp_path)
    old = _case()
    reference = await exports.prepare(user_id="owner", case=old)
    await exports.download(user_id="owner", case=old, export_id=reference["id"])
    foreign = await exports.prepare(user_id="another", case=old)
    await exports.delete_owner(user_id="owner")
    with pytest.raises(TaxExportError):
        await exports.prepare(user_id="owner", case=old)
    assert (await exports.download(user_id="another", case=old, export_id=foreign["id"])).path.exists()
    assert await exports.prepare(user_id="owner", case=_case(id="newcase", conversation_id="newconversation"))


async def test_no_saved_review_produces_no_export_link(tmp_path):
    exports = TaxReviewExports(tmp_path)
    with pytest.raises(TaxExportError) as error:
        await exports.prepare(user_id="owner", case=_case(review=None))
    assert error.value.status_code == 404 and not list(tmp_path.rglob("*.json"))


async def test_open_download_handle_survives_deletion_before_response_reads_it(tmp_path):
    exports = TaxReviewExports(tmp_path)
    case = _case()
    reference = await exports.prepare(user_id="owner", case=case)
    artifact = await exports.download(user_id="owner", case=case, export_id=reference["id"])
    await exports.delete_case(user_id="owner", case_id=case["id"])
    assert not artifact.path.exists()
    assert len(artifact.file.read()) == artifact.size
    artifact.close()
    assert artifact.file.closed


async def test_lifecycle_metadata_is_bounded_and_active_state_is_released(tmp_path):
    exports = TaxReviewExports(tmp_path, max_snapshots=2)
    for index in range(5):
        await exports.delete_case(user_id=f"owner{index}", case_id=f"case{index}")
    assert len(list(exports.root.glob("*/*/.deleted"))) == 2
    assert len(exports._owners) == 0


async def test_global_pruning_does_not_break_an_already_open_download(tmp_path):
    exports = TaxReviewExports(tmp_path, max_snapshots=1)
    case = _case()
    reference = await exports.prepare(user_id="owner", case=case)
    artifact = await exports.download(user_id="owner", case=case, export_id=reference["id"])
    await exports.prepare(user_id="another", case=case)
    assert not artifact.path.exists()
    assert artifact.file.read(4) == b"PK\x03\x04"
    artifact.close()


async def test_orphaned_outputs_from_a_killed_process_are_pruned(tmp_path):
    exports = TaxReviewExports(tmp_path)
    case = _case()
    await exports.prepare(user_id="owner", case=case)
    orphan = next(exports.root.glob("*/*/*.json")).with_name(".stopped-render.tmp")
    orphan.write_bytes(b"partial workbook")
    await exports.prepare(user_id="owner", case=case)
    assert not orphan.exists()

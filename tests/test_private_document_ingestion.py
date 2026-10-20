"""Protected document processing stays local, isolated and owner-bound."""
from __future__ import annotations

import asyncio
import json
import stat
import threading
import uuid
from types import SimpleNamespace
from unittest.mock import AsyncMock

import httpx
import pytest

from audrey.config import Config, EnvOverrides
from audrey.kb.file_deletion import FileDeletionWorker, FileOperationLocks
from audrey.kb.private_documents import (
    PrivateDocumentError,
    extract_private_document,
    private_document_source,
    private_text_collection,
    read_private_document,
)
from audrey.kb.storage_lifecycle import StorageLifecycle, StorageReservation
from audrey.kb.uploads_db import UploadsDB
from audrey.kb.user_store import sanitize_user, user_text_collection
from audrey.models.health import HealthTracker
from audrey.models.registry import ModelRegistry
from audrey.pipeline.fair_gate import FairLocalGate
from audrey.private_processing import PrivateProcessingRuntime
from audrey.processing_policy import ProcessingPolicy, ProcessingPolicyError, processing_scope
from audrey.routes import files
from audrey.routes.inflight import UserInflightRegistry
from audrey.tools.discovery import ToolRegistry

USER = "private-owner"
STAMP = "2026-10-19T12:00:00+00:00"
OLD = "2026-10-19T10:00:00+00:00"
CHAT = "local-summary:latest"
EMBED = "local-embed:latest"
PRIVATE_TEXT = "PRIVATE_DOCUMENT_EXAMPLE_2486"


@pytest.fixture
def setup(tmp_path):
    db = UploadsDB(tmp_path / "uploads.sqlite")
    cfg = Config({
        "kb": {"upload_root": str(tmp_path / "uploads"), "video": {"summarise_model": CHAT}},
        "text_embedding": {"model": EMBED, "digest": "b" * 64, "dimensions": 2,
                           "document_prefix": "", "query_prefix": "query: "},
        "model_registry": {"general": [{"name": CHAT, "location": "local"}]},
        "title_generation": {"enabled": False},
    }, EnvOverrides(_env_file=None))
    qdrant = AsyncMock()
    app = SimpleNamespace(state=SimpleNamespace(
        cfg=cfg, registry=ModelRegistry(cfg), health=HealthTracker(),
        gate=FairLocalGate(concurrency=1), inflight=UserInflightRegistry(max_inflight_per_user=3),
        tools=ToolRegistry(), qdrant=qdrant, uploads_db=db,
        file_operation_locks=FileOperationLocks(),
    ))
    seen = []
    def handle(request):
        seen.append(request)
        if request.url.path == "/api/tags":
            return httpx.Response(200, json={"models": [
                {"name": CHAT, "model": CHAT, "digest": "a" * 64},
                {"name": EMBED, "model": EMBED, "digest": "b" * 64},
            ]})
        if request.url.path == "/api/show":
            return httpx.Response(200, json={
                "details": {"format": "gguf"}, "capabilities": [],
                "model_info": {"general.architecture": "local", "general.parameter_count": 100},
            })
        if request.url.path == "/api/embed":
            payload = json.loads(request.content)
            return httpx.Response(200, json={"embeddings": [[1.0, 2.0] for _ in payload["input"]]})
        if request.url.path == "/api/chat":
            return httpx.Response(200, json={"message": {"content": "This document describes a local test. It provides a private example."}})
        pytest.fail(f"Unexpected egress path {request.url.path}")
    policy = ProcessingPolicy(local_models={CHAT: "a" * 64, EMBED: "b" * 64}, ollama_base_url="http://ollama:11434")
    runtime = PrivateProcessingRuntime(policy, ollama_transport=httpx.MockTransport(handle))
    yield SimpleNamespace(db=db, app=app, storage=StorageLifecycle(db), runtime=runtime,
                          seen=seen, qdrant=qdrant, tmp_path=tmp_path)
    db.close()


async def _land(state, suffix=".txt", content=PRIVATE_TEXT, *, private=True):
    file_id = str(uuid.uuid4())
    path = state.tmp_path / "uploads" / sanitize_user(USER) / (file_id + suffix)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(content, encoding="utf-8")
    reservation = await state.storage.reserve_single_upload(
        reservation_id=file_id, user=USER, bytes_=path.stat().st_size,
        max_user_bytes=100_000, now=STAMP, expired_before=OLD, private_required=private,
    )
    return file_id, path, reservation


async def _ingest(state, file_id, path, reservation):
    return await files._ingest_private_document(
        SimpleNamespace(app=state.app), path, runtime=state.runtime,
        storage=state.storage, reservation=reservation, user=USER,
        file_id=file_id, filename="example" + path.suffix,
        written=path.stat().st_size, max_total=100_000,
    )


async def test_private_document_roundtrip_uses_own_clients_and_collection(setup, caplog):
    state = setup
    file_id, path, reservation = await _land(state)
    result = await _ingest(state, file_id, path, reservation)
    row = await state.db.get_upload(file_id)
    assert result.status == "ready"
    assert row["private_required"] == 1
    assert stat.S_IMODE(path.with_suffix(".ocr.txt").stat().st_mode) == 0o600
    assert row["collection"] == private_text_collection(USER) != user_text_collection(USER)
    points = state.qdrant.upsert_text.call_args.args[0]
    assert all(point.payload["private_required"] is True and point.payload["user"] == USER for point in points)
    assert all(point.payload["embedding_digest"] == "b" * 64 for point in points)
    assert {r.url.host for r in state.seen} == {"ollama"}
    assert any(r.url.path == "/api/embed" and PRIVATE_TEXT.encode() in r.content for r in state.seen)
    assert PRIVATE_TEXT not in caplog.text
    with processing_scope(state.runtime.policy):
        text = await read_private_document(state.app, SimpleNamespace(storage_namespace=USER), file_id)
        assert text.text == PRIVATE_TEXT
        async with private_document_source(state.app, SimpleNamespace(storage_namespace=USER), file_id) as source:
            assert source.path == path
    with pytest.raises(ProcessingPolicyError):
        await read_private_document(state.app, SimpleNamespace(storage_namespace=USER), file_id)


@pytest.mark.parametrize("issue", ["ordinary", "forged_user", "missing", "existing"])
async def test_private_admission_refuses_before_content_calls(setup, issue):
    state = setup
    file_id, path, reservation = await _land(state, private=issue != "ordinary")
    if issue == "forged_user":
        reservation = StorageReservation(reservation.reservation_id, "other-user", "single_shot", 12, True)
    if issue == "missing":
        await state.storage.release(reservation)
    if issue == "existing":
        await state.db.record_upload(file_id=file_id, user=USER, filename="example.txt", mime="text/plain", bytes_=12,
                                    kind="text", collection="ordinary", chunks=1, uploaded_at=STAMP)
    with pytest.raises(ProcessingPolicyError):
        await _ingest(state, file_id, path, reservation)
    assert state.seen == []
    state.qdrant.upsert_text.assert_not_called()


async def test_claimed_caller_boolean_cannot_downgrade_private_reservation(setup):
    state = setup
    file_id, path, reservation = await _land(state)
    forged = StorageReservation(file_id, USER, "single_shot", reservation.reserved_bytes, False)
    await _ingest(state, file_id, path, forged)
    assert (await state.db.get_upload(file_id))["private_required"] == 1


@pytest.mark.parametrize("failure", ["embed", "upsert", "commit"])
async def test_failure_rolls_back_private_artifacts_without_disclosing_content(setup, failure, caplog):
    state = setup
    file_id, path, reservation = await _land(state)
    if failure == "embed":
        original_handler = state.runtime.ollama_transport.handler
        def fail(request):
            if request.url.path == "/api/embed":
                return httpx.Response(503, text=PRIVATE_TEXT)
            return original_handler(request)
        state.runtime = PrivateProcessingRuntime(state.runtime.policy, ollama_transport=httpx.MockTransport(fail))
    elif failure == "upsert":
        state.qdrant.upsert_text.side_effect = RuntimeError(PRIVATE_TEXT)
    else:
        state.storage.commit_upload = AsyncMock(side_effect=RuntimeError(PRIVATE_TEXT))
    with pytest.raises(ProcessingPolicyError) as error:
        await _ingest(state, file_id, path, reservation)
    assert PRIVATE_TEXT not in str(error.value) + caplog.text
    assert await state.db.get_upload(file_id) is None
    assert not path.exists()
    assert not path.with_suffix(".ocr.txt").exists()
    if failure != "embed":
        state.qdrant.delete_by_file_id.assert_awaited_once_with(file_id, user=USER, collection=private_text_collection(USER))


async def test_cancelled_commit_retains_successful_owned_document(setup):
    state = setup
    file_id, path, reservation = await _land(state)
    entered, finish = asyncio.Event(), asyncio.Event()
    commit = state.storage.commit_upload
    async def delayed(*args, **kwargs):
        entered.set()
        await finish.wait()
        await commit(*args, **kwargs)
    state.storage.commit_upload = delayed
    task = asyncio.create_task(_ingest(state, file_id, path, reservation))
    await entered.wait()
    task.cancel()
    finish.set()
    with pytest.raises(asyncio.CancelledError):
        await task
    assert (await state.db.get_upload(file_id))["private_required"] == 1
    assert path.exists() and path.with_suffix(".ocr.txt").exists()
    state.qdrant.delete_by_file_id.assert_not_called()


async def test_delete_includes_private_collection_and_derived_text(setup):
    state = setup
    file_id, path, reservation = await _land(state)
    await _ingest(state, file_id, path, reservation)
    worker = FileDeletionWorker(db=state.db, qdrant=state.qdrant, upload_root=state.tmp_path / "uploads",
                                locks=state.app.state.file_operation_locks)
    result = await worker.request(file_id, user=USER)
    assert result.completed
    assert private_text_collection(USER) in {call.kwargs["collection"] for call in state.qdrant.delete_by_file_id.call_args_list}
    assert not path.exists() and not path.with_suffix(".ocr.txt").exists()


@pytest.mark.parametrize("issue", ["other_owner", "wrong_collection", "not_ready", "symlink"])
async def test_protected_reader_refuses_stale_or_foreign_sources(setup, issue):
    state = setup
    file_id, path, reservation = await _land(state)
    await _ingest(state, file_id, path, reservation)
    owner = USER
    if issue == "other_owner":
        owner = "other-owner"
    elif issue == "wrong_collection":
        state.db._conn.execute("UPDATE uploads SET collection='ordinary' WHERE file_id=?", (file_id,))
        state.db._conn.commit()
    elif issue == "not_ready":
        state.db._conn.execute("UPDATE uploads SET status='failed' WHERE file_id=?", (file_id,))
        state.db._conn.commit()
    else:
        derived = path.with_suffix(".ocr.txt")
        derived.unlink()
        outside = state.tmp_path / "outside.txt"
        outside.write_text(PRIVATE_TEXT)
        derived.symlink_to(outside)
    with processing_scope(state.runtime.policy), pytest.raises(ProcessingPolicyError):
        await read_private_document(state.app, SimpleNamespace(storage_namespace=owner), file_id)


@pytest.mark.parametrize("suffix,content", [(".csv", "amount,total\n1,2"), (".html", "<p>Local text</p><script>ignored</script>")])
def test_supported_plain_and_html_formats_are_local(tmp_path, monkeypatch, suffix, content):
    path = tmp_path / ("example" + suffix)
    path.write_text(content)
    monkeypatch.setattr("audrey.kb.private_documents.magic.from_file", lambda *args, **kwargs: "text/html" if suffix == ".html" else "text/plain")
    mime, text = extract_private_document(path)
    assert mime in {"text/csv", "text/html"}
    assert "ignored" not in text


@pytest.mark.parametrize("issue", ["blank", "parser_error"])
def test_pdf_does_not_silently_omit_unread_pages(tmp_path, monkeypatch, issue, caplog):
    path = tmp_path / "example.pdf"
    path.write_bytes(b"%PDF-test")
    monkeypatch.setattr("audrey.kb.private_documents.magic.from_file", lambda *args, **kwargs: "application/pdf")
    class Page:
        def get_contents(self):
            return SimpleNamespace(get_data=lambda: b"painted-page")

        def extract_text(self):
            if issue == "parser_error":
                raise ValueError(PRIVATE_TEXT)
            return ""
    monkeypatch.setattr("pypdf.PdfReader", lambda *args, **kwargs: SimpleNamespace(is_encrypted=False, pages=[Page()]))
    with pytest.raises(PrivateDocumentError) as error:
        extract_private_document(path)
    assert PRIVATE_TEXT not in str(error.value) + caplog.text


async def test_pdf_summary_uses_only_policy_primary_and_embeds_both_artifacts(setup, monkeypatch):
    state = setup
    file_id, path, reservation = await _land(state, ".pdf", "%PDF-fixture")
    monkeypatch.setattr("audrey.kb.private_documents.magic.from_file", lambda *args, **kwargs: "application/pdf")
    monkeypatch.setattr("pypdf.PdfReader", lambda *args, **kwargs: SimpleNamespace(
        is_encrypted=False, pages=[SimpleNamespace(extract_text=lambda: PRIVATE_TEXT)],
    ))
    await _ingest(state, file_id, path, reservation)
    row = await state.db.get_upload(file_id)
    assert row["summary"].startswith("This document describes")
    assert path.with_suffix(".summary.txt").read_text() == row["summary"]
    assert stat.S_IMODE(path.with_suffix(".summary.txt").stat().st_mode) == 0o600
    assert "--- Page 1 ---" in path.with_suffix(".ocr.txt").read_text()
    points = state.qdrant.upsert_text.call_args.args[0]
    assert {point.payload["artifact"] for point in points} == {"text", "summary"}
    assert {json.loads(request.content)["model"] for request in state.seen if request.url.path == "/api/chat"} == {CHAT}


async def test_unapproved_summary_primary_is_skipped_without_external_fallback(setup, monkeypatch):
    state = setup
    state.app.state.cfg.raw["kb"]["video"]["summarise_model"] = "external:cloud"
    file_id, path, reservation = await _land(state, ".pdf", "%PDF-fixture")
    monkeypatch.setattr("audrey.kb.private_documents.magic.from_file", lambda *args, **kwargs: "application/pdf")
    monkeypatch.setattr("pypdf.PdfReader", lambda *args, **kwargs: SimpleNamespace(
        is_encrypted=False, pages=[SimpleNamespace(extract_text=lambda: PRIVATE_TEXT)],
    ))
    await _ingest(state, file_id, path, reservation)
    assert (await state.db.get_upload(file_id))["summary"] == ""
    assert not any(request.url.path == "/api/chat" for request in state.seen)


def test_xlsx_reads_all_sheets_without_evaluating_formulas(tmp_path):
    from openpyxl import Workbook
    path = tmp_path / "example.xlsx"
    book = Workbook()
    book.active.append(["amount", "formula"])
    book.active.append([10, "=SUM(A2:A2)"])
    book.create_sheet("Second").append(["Second-sheet data"])
    book.save(path)
    book.close()
    mime, text = extract_private_document(path)
    assert "spreadsheetml" in mime
    assert "Second-sheet data" in text and "=SUM(A2:A2)" in text


def test_docx_includes_paragraphs_and_tables(tmp_path):
    from docx import Document
    path = tmp_path / "example.docx"
    doc = Document()
    doc.add_paragraph("Main private document")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Table evidence"
    doc.save(path)
    _, text = extract_private_document(path)
    assert "Main private document" in text and "Table evidence" in text


async def test_cancelled_derived_write_joins_thread_before_unlink(setup, monkeypatch):
    state = setup
    file_id, path, reservation = await _land(state)
    entered, finish = threading.Event(), threading.Event()
    write = files._write_private_text
    def delayed_write(*args):
        entered.set()
        assert finish.wait(5)
        write(*args)
    monkeypatch.setattr(files, "_write_private_text", delayed_write)
    task = asyncio.create_task(_ingest(state, file_id, path, reservation))
    assert await asyncio.to_thread(entered.wait, 5)
    task.cancel()
    await asyncio.sleep(0)
    assert not task.done()
    finish.set()
    with pytest.raises(asyncio.CancelledError):
        await task
    assert await state.db.get_upload(file_id) is None
    assert not path.exists() and not path.with_suffix(".ocr.txt").exists()
    state.qdrant.upsert_text.assert_not_called()


async def test_pending_private_deletion_is_not_readable_and_redacts_errors(setup, caplog):
    state = setup
    file_id, path, reservation = await _land(state)
    await _ingest(state, file_id, path, reservation)
    state.qdrant.delete_by_file_id.side_effect = RuntimeError(PRIVATE_TEXT)
    worker = FileDeletionWorker(db=state.db, qdrant=state.qdrant, upload_root=state.tmp_path / "uploads",
                                locks=state.app.state.file_operation_locks)
    result = await worker.request(file_id, user=USER)
    assert not result.completed
    assert result.error == "Private file cleanup is pending."
    assert PRIVATE_TEXT not in caplog.text
    with processing_scope(state.runtime.policy), pytest.raises(ProcessingPolicyError):
        await read_private_document(state.app, SimpleNamespace(storage_namespace=USER), file_id)
    stats = await state.db.pending_file_deletions()
    assert all(PRIVATE_TEXT not in str(row) for row in stats)


def _pdf_with_pages(path, kinds):
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    for kind in kinds:
        page = writer.add_blank_page(width=200, height=100)
        if kind == "text":
            page[NameObject("/Resources")] = DictionaryObject({
                NameObject("/Font"): DictionaryObject({
                    NameObject("/F1"): DictionaryObject({
                        NameObject("/Type"): NameObject("/Font"),
                        NameObject("/Subtype"): NameObject("/Type1"),
                        NameObject("/BaseFont"): NameObject("/Helvetica"),
                    }),
                }),
            })
            data = b"BT /F1 12 Tf 20 40 Td (Private document evidence) Tj ET"
        elif kind == "painted":
            data = b"0 0 10 10 re f"
        elif kind == "empty_stream":
            data = b" \n\t"
        else:
            continue
        stream = DecodedStreamObject()
        stream.set_data(data)
        page[NameObject("/Contents")] = writer._add_object(stream)
    writer.write(path)
    writer.close()


@pytest.mark.parametrize("blank_kind", ["no_stream", "empty_stream"])
def test_pdf_keeps_provably_blank_page_numbers(tmp_path, blank_kind):
    path = tmp_path / "text-and-blank.pdf"
    _pdf_with_pages(path, ["text", blank_kind, "text"])
    mime, text = extract_private_document(path)
    assert mime == "application/pdf"
    assert "--- Page 1 ---" in text and "--- Page 2 ---" in text and "--- Page 3 ---" in text
    assert text.count("Private document evidence") == 2


def test_entirely_blank_pdf_is_refused(tmp_path):
    path = tmp_path / "all-blank.pdf"
    _pdf_with_pages(path, ["no_stream", "empty_stream"])
    with pytest.raises(PrivateDocumentError) as error:
        extract_private_document(path)
    assert error.value.code == "private_document_invalid"


def test_pdf_painted_page_without_text_still_requires_ocr(tmp_path):
    path = tmp_path / "painted.pdf"
    _pdf_with_pages(path, ["text", "painted"])
    with pytest.raises(PrivateDocumentError) as error:
        extract_private_document(path)
    assert error.value.code == "private_document_requires_ocr"

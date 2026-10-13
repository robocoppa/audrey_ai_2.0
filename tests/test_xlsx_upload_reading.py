"""XLSX uploads enter the readable-file contract used by chat and My Files."""
from __future__ import annotations

import zipfile

from openpyxl import Workbook

from audrey.kb import extract


def test_real_excel_source_is_allowed_and_readable(tmp_path):
    path = tmp_path / "costs.xlsx"
    workbook = Workbook()
    workbook.active.title = "Costs"
    workbook.active.append(["Project", "Amount"])
    workbook.active.append(["Alder", 12.50])
    workbook.create_sheet("Notes").append(["Other sheet"])
    workbook.save(path)
    workbook.close()

    assert ".xlsx" in extract.ALLOWED_EXTENSIONS
    mime = extract.sniff_mime(path)
    assert mime in extract.ALLOWED_TEXT_MIMES
    assert mime == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    text = extract.extract_uploaded_text(path)
    assert "Costs" in text and "Alder" in text and "12.5" in text
    assert "preview" in text.lower() and "Notes" in text
    assert "Preview only" in text


def test_generic_zip_named_excel_does_not_gain_spreadsheet_mime(tmp_path):
    path = tmp_path / "fake.xlsx"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("anything.txt", "not a spreadsheet")
    assert extract.sniff_mime(path) not in extract.ALLOWED_MIMES


def test_blank_first_tab_does_not_hide_later_uploaded_data(tmp_path):
    from audrey.kb.table_analysis import TableAnalysisRequest, analyze_table
    path = tmp_path / "blank-cover.xlsx"
    workbook = Workbook()
    workbook.active.title = "Blank"
    data = workbook.create_sheet("Costs")
    data.append(["Amount"])
    data.append([12.5])
    workbook.save(path)
    workbook.close()
    mime = extract.sniff_mime(path)
    assert "12.5" in extract.extract_uploaded_text(path)
    result = analyze_table(path, mime, TableAnalysisRequest())
    assert result["sheet"] == "Costs"
    assert result["sheets"] == ["Blank", "Costs"]

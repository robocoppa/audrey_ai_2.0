"""Supported formatting is projected from saved bytes without changing content."""
from __future__ import annotations

import io
import zipfile

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from audrey.file_previews import FilePreviewError, preview_generated_file, workbook_preview


def _canonical(document):
    values = []
    for item in document.iter_inner_content():
        if isinstance(item, Paragraph):
            values.append(item.text)
        elif isinstance(item, Table):
            values.append('\n'.join('\t'.join(cell.text for cell in row.cells) for row in item.rows))
    return '\n\n'.join(values)


def _saved(tmp_path, document, text=None):
    source = tmp_path / 'report.docx'
    sidecar = tmp_path / 'report.ocr.txt'
    document.save(source)
    sidecar.write_text(_canonical(document) if text is None else text)
    return source, sidecar


def test_docx_uses_actual_headings_emphasis_lists_tables_and_code(tmp_path):
    document = Document()
    document.add_heading('Review summary', level=1)
    paragraph = document.add_paragraph()
    paragraph.add_run('Important ').bold = True
    paragraph.add_run('supporting text').italic = True
    document.add_paragraph('First task', 'List Bullet')
    document.add_paragraph('Next task', 'List Number')
    document.add_paragraph('A quotation', 'Quote')
    document.styles.add_style('CodeBlock', WD_STYLE_TYPE.PARAGRAPH)
    document.styles.add_style('InlineCode', WD_STYLE_TYPE.CHARACTER)
    document.add_paragraph('print("literal")', 'CodeBlock')
    paragraph = document.add_paragraph('Use ')
    paragraph.add_run('<script>not executed</script>', 'InlineCode')
    table = document.add_table(rows=2, cols=2, style='Table Grid')
    for index, value in enumerate(['Item', 'Amount', 'Meals', '$25']):
        cell = table.cell(index // 2, index % 2)
        cell.text = value
        if index < 2:
            cell.paragraphs[0].runs[0].bold = True
    source, sidecar = _saved(tmp_path, document)
    original = source.read_bytes()
    result = preview_generated_file(source, format='docx', text_path=sidecar)
    blocks = result['blocks']
    assert result['text'] == sidecar.read_text()
    assert result['formatting_truncated'] is False
    assert blocks[0]['kind'] == 'heading' and blocks[0]['level'] == 1
    assert blocks[0]['runs'][0]['text'] == 'Review summary'
    assert blocks[1]['runs'] == [{'text': 'Important ', 'bold': True}, {'text': 'supporting text', 'italic': True}]
    assert [block.get('list') for block in blocks[2:4]] == ['bullet', 'number']
    assert blocks[4]['kind'] == 'quote' and blocks[5]['kind'] == 'code'
    assert blocks[6]['runs'][1] == {'text': '<script>not executed</script>', 'code': True}
    assert blocks[7]['kind'] == 'table' and blocks[7]['header'] is True
    assert blocks[7]['rows'][1][0]['runs'] == [{'text': 'Meals'}]
    assert source.read_bytes() == original


def test_docx_paging_slices_styled_runs_without_repetition(tmp_path):
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run('A' * 99_990).bold = True
    paragraph.add_run('B' * 30).italic = True
    source, sidecar = _saved(tmp_path, document)
    first = preview_generated_file(source, format='docx', text_path=sidecar)
    second = preview_generated_file(source, format='docx', text_path=sidecar, text_offset=first['next_offset'])
    assert first['blocks'][0]['runs'] == [
        {'text': 'A' * 99_990, 'bold': True}, {'text': 'B' * 10, 'italic': True},
    ]
    assert second['blocks'][0]['runs'] == [{'text': 'B' * 20, 'italic': True}]
    assert first['text'] + second['text'] == sidecar.read_text()
    assert not first['formatting_truncated'] and not second['formatting_truncated']


def test_docx_unmappable_or_oversized_table_falls_back_without_losing_text(tmp_path):
    document = Document()
    document.add_heading('Actual source', 1)
    source, sidecar = _saved(tmp_path, document, text='Different saved canonical text')
    result = preview_generated_file(source, format='docx', text_path=sidecar)
    assert result['blocks'] == [] and result['formatting_truncated']
    assert result['text'] == 'Different saved canonical text'
    table = document.add_table(rows=201, cols=1)
    for row in table.rows:
        row.cells[0].text = 'Retained content'
    source, sidecar = _saved(tmp_path, document)
    result = preview_generated_file(source, format='docx', text_path=sidecar)
    assert result['blocks'] == [] and result['formatting_truncated']
    assert result['text'] == sidecar.read_text()


def test_partial_table_is_text_on_each_section_with_no_omitted_cells(tmp_path):
    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'A' * 99_990
    table.cell(0, 1).text = 'B' * 50
    source, sidecar = _saved(tmp_path, document)
    first = preview_generated_file(source, format='docx', text_path=sidecar)
    second = preview_generated_file(source, format='docx', text_path=sidecar, text_offset=first['next_offset'])
    assert first['formatting_truncated'] and second['formatting_truncated']
    assert first['blocks'][0]['runs'][0]['text'] + second['blocks'][0]['runs'][0]['text'] == sidecar.read_text()


def test_docx_checks_zip_inflation_before_parsing(tmp_path):
    source, sidecar = tmp_path / 'bad.docx', tmp_path / 'bad.ocr.txt'
    sidecar.write_text('safe canonical text')
    with zipfile.ZipFile(source, 'w', compression=zipfile.ZIP_DEFLATED) as package:
        package.writestr('word/document.xml', b'x' * (17 * 1024 * 1024))
    with pytest.raises(FilePreviewError, match='large'):
        preview_generated_file(source, format='docx', text_path=sidecar)


def test_workbook_display_formats_styles_widths_and_literal_formula_are_from_source():
    book = Workbook()
    sheet = book.active
    sheet.append(['Amount', 'Count', 'Ratio', 'Literal'])
    sheet.append([-1234.505, 12000, 0.125, '=NOT_EXECUTED()'])
    formats = ['$#,##0.00', '#,##0', '0.00%', 'General']
    for index, number_format in enumerate(formats, 1):
        sheet.cell(2, index).number_format = number_format
        sheet.cell(1, index).fill = PatternFill('solid', fgColor='16324F')
        sheet.cell(1, index).font = Font(bold=True, color='FFFFFF')
    sheet.cell(2, 1).font = Font(italic=True)
    sheet.cell(2, 1).fill = PatternFill('solid', fgColor='F2F6FA')
    sheet.column_dimensions['A'].width = 20
    sheet.column_dimensions['B'].width = 100
    data = io.BytesIO()
    book.save(data)
    data.seek(0)
    result = workbook_preview(data)['sheets'][0]
    assert result['rows'][1] == [-1234.505, 12000, 0.125, '=NOT_EXECUTED()']
    assert result['column_widths'] == [20, 60]
    assert all(style == {'bold': True, 'header': True} for style in result['cell_styles'][0])
    assert result['cell_styles'][1] == [
        {'italic': True, 'banded': True, 'display': '-$1,234.51'},
        {'display': '12,000'}, {'display': '12.50%'}, {},
    ]
    assert all('align' not in style for row in result['cell_styles'] for style in row)


def test_workbook_paging_keeps_styles_at_their_original_rows_and_does_not_interpret_formats():
    book = Workbook()
    sheet = book.active
    for index in range(205):
        sheet.append([index, True, '<img src="https://invalid/track">'])
        sheet.cell(index + 1, 1).number_format = '#,##0.00'
    sheet.cell(201, 1).number_format = '[Blue]"<script>format</script>"0'
    sheet.cell(202, 1).font = Font(bold=True)
    sheet.cell(202, 1).fill = PatternFill('solid', fgColor='FF0000')
    data = io.BytesIO()
    book.save(data)
    data.seek(0)
    page = workbook_preview(data, sheet=0, row_offset=200)['sheets'][0]
    assert page['rows'][0] == [200, True, '<img src="https://invalid/track">']
    assert page['cell_styles'][0] == [{}, {}, {}]
    assert page['cell_styles'][1][0] == {'bold': True, 'display': '201.00'}
    assert len(page['rows']) == len(page['cell_styles']) == 5

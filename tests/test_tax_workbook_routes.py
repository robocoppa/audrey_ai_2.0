"""Saved-case exports retain exact sources and native delivery without model work."""
from __future__ import annotations

from dataclasses import replace
from io import BytesIO

import pytest
from openpyxl import load_workbook
from test_native_private_reviews import case_tool, saved_case, start_review
from test_native_private_workflow import _settle
from test_native_private_workflow import document_setup as document_setup
from test_native_private_workflow import native as native
from test_tax_review import _call, _create, _evidence
from test_tax_review import state as state

from audrey.tax_exports import TaxExportError, TaxReviewExports


async def _saved(state, tmp_path):
    state.app.state.tax_review_exports = TaxReviewExports(tmp_path / 'exports')
    case = await _create(state)
    evidence = await _call(state, 'record_evidence', expected_revision=case['revision'], data=_evidence(state))
    assert evidence['persistence']['status'] == 'saved'
    saved = await _call(state, 'save_review', expected_revision=evidence['revision'], data={
        'summary': 'A retained official passage, with applicability unresolved.',
        'findings': [{'label': 'Retained source', 'detail': 'Dated official supporting passage.',
                      'evidence_ids': [evidence['updated_entry']['id']]}],
        'next_question': '',
    })
    return saved


async def test_prepare_is_lazy_and_download_contains_retained_source_and_citation(state, tmp_path):
    saved = await _saved(state, tmp_path)
    response = await state.client.post(state.path + '/exports', json={'expected_revision': saved['revision']})
    assert response.status_code == 200, response.text
    reference = response.json()['export']
    assert list((tmp_path / 'exports').rglob('*.xlsx')) == []
    download = await state.client.get(reference['download_url'])
    assert download.status_code == 200, download.text if download.status_code != 200 else ''
    assert 'attachment' in download.headers['content-disposition']
    assert download.headers['cache-control'] == 'private, no-store'
    assert download.headers['x-content-type-options'] == 'nosniff'
    assert download.headers['x-tax-case-revision'] == str(saved['revision'])
    workbook = load_workbook(BytesIO(download.content), read_only=True, data_only=False)
    try:
        assert workbook.sheetnames == ['Overview', 'Facts', 'Evidence', 'Calculations', 'Unresolved items']
        evidence = '\n'.join(str(cell.value) for row in workbook['Evidence'] for cell in row)
        assert 'Dated official supporting passage.' in evidence
        assert state.metadata['source_url'] in evidence
        assert state.metadata['revision_sha256'] in evidence
        assert 'proposed' in evidence
        evidence_rows = list(workbook['Evidence'].values)
        assert dict(zip(evidence_rows[0], evidence_rows[1], strict=True))['Revision'] == state.metadata['revision_date']
    finally:
        workbook.close()
    repeated = await state.client.get(reference['download_url'])
    assert repeated.content == download.content
    current = (await state.client.get(state.path)).json()['case']
    assert current['revision'] == saved['revision']
    assert current['evidence'][0]['passage'] == 'Dated official supporting passage.'


async def test_old_link_stays_at_old_revision_and_stale_prepare_is_not_silent(state, tmp_path):
    saved = await _saved(state, tmp_path)
    reference = (await state.client.post(state.path + '/exports', json={'expected_revision': saved['revision']})).json()['export']
    changed = await state.client.post(state.path + '/actions', json={
        'operation': 'change_year', 'expected_revision': saved['revision'], 'data': {'tax_year': 2024},
    })
    assert changed.status_code == 200, changed.text
    assert (await state.client.post(state.path + '/exports', json={'expected_revision': saved['revision']})).status_code == 409
    download = await state.client.get(reference['download_url'])
    assert download.status_code == 200
    workbook = load_workbook(BytesIO(download.content), read_only=True)
    try:
        overview = dict(workbook['Overview'].values)
        assert overview['Tax year'] == '2025'
        assert overview['Case revision'] == str(saved['revision'])
    finally:
        workbook.close()


async def test_exports_require_owned_private_case_and_full_scope(state, tmp_path):
    saved = await _saved(state, tmp_path)
    reference = (await state.client.post(state.path + '/exports', json={'expected_revision': saved['revision']})).json()['export']
    state.current[0] = replace(state.owner, auth_method='personal_token', scopes=frozenset({'account:read'}))
    assert (await state.client.get(reference['download_url'])).status_code == 403
    state.current[0] = await state.store.resolve_external_identity(
        provider='test', subject='export-other', email='other@example.com', display_name='Other',
        role='user', auth_method='test',
    )
    assert (await state.client.get(reference['download_url'])).status_code == 404
    assert (await state.client.post(state.path + '/exports', json={'expected_revision': saved['revision']})).status_code == 404
    state.current[0] = state.owner
    await state.store.conversations.delete(user_id=state.owner.user_id, conversation_id=state.conversation.conversation_id)
    assert (await state.client.get(reference['download_url'])).status_code == 404


async def test_export_failure_keeps_saved_evidence_and_worksheet(state, tmp_path, monkeypatch):
    saved = await _saved(state, tmp_path)
    async def unavailable(**_kwargs):
        raise TaxExportError('tax_export_unavailable', 503, 'Workbook preparation is unavailable.')
    monkeypatch.setattr(state.app.state.tax_review_exports, 'prepare', unavailable)
    response = await state.client.post(state.path + '/exports', json={'expected_revision': saved['revision']})
    assert response.status_code == 503
    current = (await state.client.get(state.path)).json()['case']
    assert current['revision'] == saved['revision']
    assert current['review']['summary'] == 'A retained official passage, with applicability unresolved.'
    assert current['evidence'][0]['passage'] == 'Dated official supporting passage.'


async def test_export_requires_saved_review_and_does_not_require_live_reference_services(state, tmp_path):
    state.app.state.tax_review_exports = TaxReviewExports(tmp_path / 'exports')
    case = await _create(state)
    assert (await state.client.post(state.path + '/exports', json={'expected_revision': case['revision']})).status_code == 409
    saved = await _call(state, 'save_review', expected_revision=case['revision'], data={
        'summary': 'No supporting passage recorded yet.', 'findings': [], 'next_question': '',
    })
    state.app.state.reference_releases = None
    state.app.state.reference_libraries = None
    response = await state.client.post(state.path + '/exports', json={'expected_revision': saved['revision']})
    assert response.status_code == 200, response.text
    assert (await state.client.get(response.json()['export']['download_url'])).status_code == 200


@pytest.mark.parametrize('export_unavailable', [False, True])
async def test_native_saved_answer_adds_workbook_link_without_another_model_call(native, tmp_path, monkeypatch, export_unavailable):
    state = native
    conversation, case = await saved_case(state)
    exports = TaxReviewExports(tmp_path / 'exports')
    state.app.state.tax_review_exports = exports
    if export_unavailable:
        async def unavailable(**_kwargs):
            raise TaxExportError('tax_export_unavailable', 503, 'Workbook preparation is unavailable.')
        monkeypatch.setattr(exports, 'prepare', unavailable)
    calls = []
    async def stream(execution_app, _payload, _messages, _options, *, event_context, model_tools, **_kwargs):
        calls.append(event_context.run_id)
        result = await case_tool(state, execution_app, model_tools, 'save_review',
            expected_revision=case['revision'], data={
                'summary': 'The selected review remains partial.', 'findings': [], 'next_question': '',
            })
        assert result['persistence']['status'] == 'saved'
        emitter = event_context.emitter
        emitter.run_started()
        emitter.message_started()
        emitter.text_delta('The partial review is saved.')
        emitter.message_finished(status='completed')
        emitter.run_finished(status='succeeded', finish_reason='stop')
        yield 'terminal'
    state.manager._stream_factory = stream
    run = await start_review(state, conversation['id'], 'Save the useful partial review.')
    await _settle(state, run['id'])
    messages = await state.store.conversations.list_messages(user_id=state.owner.user_id, conversation_id=conversation['id'])
    assert len(messages) == 2
    content = messages[-1].content
    assert content.startswith('The partial review is saved.')
    assert ('Download review workbook (.xlsx)' in content) is not export_unavailable
    assert calls == [run['id']]
    record = await state.store.conversations.get_run(user_id=state.owner.user_id, run_id=run['id'])
    assert record.status == 'succeeded'
    assert await state.store.private_tasks.pending_deliveries() == []
